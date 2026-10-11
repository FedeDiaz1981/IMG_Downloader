from __future__ import annotations

import os
import base64
import csv
import json
import re
import shutil
import subprocess
import threading
import time
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from difflib import SequenceMatcher
from io import BytesIO
from pathlib import Path
from typing import Iterable
from urllib.parse import quote_plus, unquote_plus, urljoin, urlparse, urlunparse

import requests
from bs4 import BeautifulSoup
from flask import Flask, jsonify, redirect, render_template_string, request, send_file, url_for
from openpyxl import load_workbook
from PIL import Image, ImageOps


BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
DOWNLOAD_DIR = BASE_DIR / "downloads"
JOB_DIR = BASE_DIR / "jobs"
BROWSER_PROFILE_DIR = BASE_DIR / ".pintofruta-edge"
PINFOFRUTA_ADMIN_URL = "https://pintofruta.com/admin"
PINFOFRUTA_DEBUG_URL = "http://127.0.0.1:9222"
PINFOFRUTA_ENV_DEFAULT = (
    Path(os.environ["PINFOFRUTA_ENV_FILE"]).expanduser()
    if os.environ.get("PINFOFRUTA_ENV_FILE")
    else BASE_DIR.parent.parent / "Proyectos" / "en curso" / "Pintofruta" / "PintoFruta_Store" / ".env.local"
)
PINFOFRUTA_INSTAGRAM_URL = "https://www.instagram.com/distribuidora_pintofruta/"
PINFOFRUTA_EDGE_CANDIDATES = (
    Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
    Path(r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"),
)

for folder in (UPLOAD_DIR, DOWNLOAD_DIR, JOB_DIR):
    folder.mkdir(exist_ok=True)


app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024

JOBS: dict[str, "JobStatus"] = {}
JOBS_LOCK = threading.Lock()
FOLDER_DIALOG_LOCK = threading.Lock()

STOPWORDS = {
    "de",
    "del",
    "la",
    "el",
    "los",
    "las",
    "para",
    "con",
    "sin",
    "por",
    "y",
    "o",
    "a",
    "en",
    "un",
    "una",
    "x",
    "mm",
    "cm",
    "kg",
    "gr",
    "ml",
    "lt",
    "mts",
    "color",
    "modelo",
}

IMAGE_EXTENSIONS = (".jpg", ".jpeg", ".png", ".webp", ".gif", ".bmp")
SKIP_IMAGE_HINTS = (
    "logo",
    "icon",
    "sprite",
    "favicon",
    "placeholder",
    "loading",
    "banner",
    "payment",
    "social",
)


def get_config_value(name: str) -> str:
    """Read configuration from the process or the Windows user environment."""
    value = os.environ.get(name, "").strip()
    if value or os.name != "nt":
        return value

    try:
        import winreg

        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, "Environment") as key:
            stored_value, _ = winreg.QueryValueEx(key, name)
        return str(stored_value).strip()
    except (FileNotFoundError, OSError):
        return ""


@dataclass
class Product:
    row: int
    brand: str
    name: str
    code: str = ""


@dataclass
class ImageCandidate:
    url: str
    text: str
    page_url: str
    score: int = 0


@dataclass
class AiVerdict:
    accepted: bool | None
    confidence: int = 0
    reason: str = ""


@dataclass
class AiCardSelection:
    selected_index: int | None
    confidence: int = 0
    reason: str = ""


@dataclass
class ManualCandidate:
    title: str
    filter_query: str
    article_index: int


@dataclass
class ManualChoice:
    id: str
    product_name: str
    brand: str
    source_image: Path
    candidates: list[ManualCandidate]


@dataclass
class PageSnapshot:
    url: str
    title: str
    text: str
    links: list[tuple[str, str]] = field(default_factory=list)
    images: list[ImageCandidate] = field(default_factory=list)


@dataclass
class ExcelLayout:
    header_row: int
    headers: list[str]


@dataclass
class JobStatus:
    id: str
    state: str = "queued"
    progress: int = 0
    message: str = "Esperando..."
    log: list[str] = field(default_factory=list)
    zip_path: Path | None = None
    output_dir: Path | None = None
    error: str | None = None
    missing_searches: list[tuple[str, str]] = field(default_factory=list)
    manual_choices: list[ManualChoice] = field(default_factory=list)

    def add(self, message: str) -> None:
        timestamp = datetime.now().strftime("%H:%M:%S")
        self.log.append(f"[{timestamp}] {message}")
        self.log = self.log[-300:]
        self.message = message


def normalize_text(value: object) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii").lower()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def slugify(value: str, fallback: str = "item") -> str:
    text = normalize_text(value)
    text = re.sub(r"\s+", "-", text).strip("-")
    return text[:90] or fallback


def readable_slug(value: str) -> str:
    return re.sub(r"\s+", " ", value.replace("-", " ")).strip()


def product_tokens(product: Product) -> set[str]:
    raw = normalize_text(f"{product.brand} {product.name} {product.code}")
    return {token for token in raw.split() if len(token) > 1 and token not in STOPWORDS}


def product_match_tokens(product: Product) -> set[str]:
    brand_tokens = set(normalize_text(product.brand).split())
    name_tokens = set(product_identity_tokens(product.name, product.brand))
    specific = {
        token
        for token in name_tokens - brand_tokens
        if len(token) > 2 and token not in STOPWORDS
    }
    return specific or brand_tokens


def safe_url(value: str) -> str:
    value = value.strip()
    if not value:
        return ""
    if not re.match(r"^https?://", value, flags=re.I):
        value = f"https://{value}"
    return value


def same_domain(base_url: str, candidate_url: str) -> bool:
    base = urlparse(base_url)
    candidate = urlparse(candidate_url)
    if not candidate.netloc:
        return True
    return candidate.netloc.lower().removeprefix("www.") == base.netloc.lower().removeprefix("www.")


def clean_url(url: str) -> str:
    parsed = urlparse(url)
    return urlunparse((parsed.scheme, parsed.netloc, parsed.path, "", parsed.query, ""))


def high_resolution_variants(url: str) -> list[str]:
    parsed = urlparse(url)
    is_tiendanube_image = (
        ("mitiendanube.com" in parsed.netloc.lower() or "cloudfront.net" in parsed.netloc.lower())
        and "/stores/" in parsed.path
        and re.search(r"-\d{2,4}-0\.[a-z0-9]{2,5}$", parsed.path, flags=re.I) is not None
    )
    if not is_tiendanube_image:
        return []

    variants: list[str] = []
    for size in (1024, 640, 480, 320):
        candidate_path = re.sub(
            r"-\d{2,4}-0(?=\.[a-z0-9]{2,5}$)",
            f"-{size}-0",
            parsed.path,
            flags=re.I,
        )
        variants.append(urlunparse((parsed.scheme, parsed.netloc, candidate_path, "", "", "")))

    # Tiendanube may expose the 1024px WebP as JPEG instead.
    if parsed.path.lower().endswith(".webp"):
        jpeg_path = re.sub(r"-\d{2,4}-0\.webp$", "-1024-1024.jpg", parsed.path, flags=re.I)
        variants.insert(0, urlunparse((parsed.scheme, parsed.netloc, jpeg_path, "", "", "")))
    return variants


def likely_header_score(values: Iterable[object]) -> int:
    cells = [normalize_text(value) for value in values if str(value or "").strip()]
    joined = " ".join(cells)
    score = 0
    if "marca" in cells or "brand" in cells:
        score += 30
    if any(term in joined for term in ("producto", "detalle", "descripcion", "variedad", "item")):
        score += 30
    if any(term in joined for term in ("presentacion", "precio", "subtotal", "cantidad")):
        score += 6
    if len(cells) >= 3:
        score += min(len(cells), 8)
    return score


def read_excel_layout(path: Path) -> ExcelLayout:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return ExcelLayout(header_row=1, headers=[])

    best_index = 0
    best_score = -1
    for index, row in enumerate(rows[:80]):
        score = likely_header_score(row)
        if score > best_score:
            best_index = index
            best_score = score

    if best_score <= 0:
        for index, row in enumerate(rows):
            if any(str(cell or "").strip() for cell in row):
                best_index = index
                break

    headers = [str(cell or "").strip() for cell in rows[best_index] if str(cell or "").strip()]
    return ExcelLayout(header_row=best_index + 1, headers=headers)


def parse_excel(
    path: Path,
    header_row: int,
    brand_column: str,
    product_column: str,
    code_column: str = "",
    extra_column: str = "",
) -> list[Product]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows or header_row < 1 or header_row > len(rows):
        return []

    headers = [str(cell or "").strip() for cell in rows[header_row - 1]]
    index = {header: idx for idx, header in enumerate(headers)}
    products: list[Product] = []

    for row_number, row in enumerate(rows[header_row:], start=header_row + 1):
        brand = str(row[index[brand_column]] or "").strip() if brand_column in index else ""
        name = str(row[index[product_column]] or "").strip() if product_column in index else ""
        code = str(row[index[code_column]] or "").strip() if code_column and code_column in index else ""
        extra = str(row[index[extra_column]] or "").strip() if extra_column and extra_column in index else ""
        if extra and normalize_text(extra) not in normalize_text(name):
            name = f"{name} {extra}".strip()
        if brand and name:
            products.append(Product(row=row_number, brand=brand, name=name, code=code))

    return products


def grouped_brands(products: Iterable[Product]) -> dict[str, list[Product]]:
    groups: dict[str, list[Product]] = {}
    for product in products:
        groups.setdefault(product.brand.strip(), []).append(product)
    return dict(sorted(groups.items(), key=lambda item: normalize_text(item[0])))


def make_session() -> requests.Session:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/125.0 Safari/537.36 IMG-Downloader/1.0"
            ),
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
        }
    )
    return session


def fetch_page(session: requests.Session, url: str, timeout: int = 15) -> str:
    response = session.get(url, timeout=timeout)
    response.raise_for_status()
    content_type = response.headers.get("content-type", "")
    if "text" not in content_type and "html" not in content_type and content_type:
        return ""
    return response.text


def image_urls_from_tag(tag, page_url: str) -> list[str]:
    attrs = [
        "data-lazy-src",
        "data-zoom-image",
        "data-large-image",
        "data-full",
        "data-original",
        "data-src",
        "srcset",
        "data-srcset",
        "src",
    ]
    urls: list[str] = []
    for attr in attrs:
        value = tag.get(attr)
        if not value:
            continue
        if attr.endswith("srcset"):
            entries: list[tuple[float, str]] = []
            for entry in str(value).split(","):
                parts = entry.strip().split()
                if not parts:
                    continue
                descriptor = parts[1] if len(parts) > 1 else "0"
                multiplier = float(descriptor[:-1]) if descriptor.endswith("x") else 0
                width = float(descriptor[:-1]) if descriptor.endswith("w") else 0
                entries.append((max(width, multiplier), parts[0]))
            urls.extend(urljoin(page_url, item[1]) for item in sorted(entries, reverse=True))
        else:
            urls.append(urljoin(page_url, str(value).strip()))

    parent_link = tag.find_parent("a", href=True)
    if parent_link:
        href = urljoin(page_url, parent_link.get("href", ""))
        if href.lower().split("?")[0].endswith(IMAGE_EXTENSIONS):
            urls.insert(0, href)

    unique_urls: list[str] = []
    seen: set[str] = set()
    for image_url in urls:
        image_url = clean_url(image_url)
        candidates = high_resolution_variants(image_url) + [image_url]

        # Many stores publish thumbnails as product-300x300.jpg while the
        # original is available at product.jpg.
        original_url = re.sub(
            r"[-_]\d{2,4}x\d{2,4}(?=\.[a-z0-9]{2,5}(?:$|\?))",
            "",
            image_url,
            flags=re.I,
        )
        if original_url != image_url:
            candidates.insert(0, original_url)

        for candidate_url in candidates:
            if not candidate_url or candidate_url in seen:
                continue
            seen.add(candidate_url)
            unique_urls.append(candidate_url)

    return unique_urls


def snapshot_page(session: requests.Session, url: str) -> PageSnapshot | None:
    try:
        html = fetch_page(session, url)
    except requests.RequestException:
        return None

    if not html:
        return None

    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(" ", strip=True) if soup.title else ""
    text = soup.get_text(" ", strip=True)
    links: list[tuple[str, str]] = []
    images: list[ImageCandidate] = []

    for meta in soup.select("meta[property='og:image'], meta[name='twitter:image']"):
        image_url = meta.get("content")
        if image_url:
            images.append(ImageCandidate(url=urljoin(url, image_url), text=f"{title} og image", page_url=url))

    for link in soup.find_all("a", href=True):
        href = urljoin(url, link.get("href", ""))
        href = clean_url(href)
        if href.startswith("http") and same_domain(url, href):
            label = link.get_text(" ", strip=True)
            links.append((href, label))

    for img in soup.find_all("img"):
        image_urls = image_urls_from_tag(img, url)
        if not image_urls:
            continue
        alt = img.get("alt", "")
        title_attr = img.get("title", "")
        surrounding = img.parent.get_text(" ", strip=True) if img.parent else ""
        for image_url in image_urls:
            text_bits = [title, alt, title_attr, surrounding, image_url]
            images.append(ImageCandidate(url=image_url, text=" ".join(text_bits), page_url=url))

    return PageSnapshot(url=url, title=title, text=text, links=links, images=images)


def score_text_against_tokens(text: str, tokens: set[str]) -> int:
    normalized = normalize_text(text)
    if not normalized:
        return 0

    score = 0
    words = set(normalized.split())
    for token in tokens:
        if token in words:
            score += 4
        elif len(token) >= 4 and token in normalized:
            score += 2
    return score


def image_quality_score(url: str) -> int:
    normalized = normalize_text(url)
    raw_url = url.lower()
    score = 0
    if any(term in normalized for term in ("original", "full", "zoom", "large", "master", "hires", "highres")):
        score += 7
    if any(term in normalized for term in ("thumb", "thumbnail", "small", "tiny", "mini")):
        score -= 8
    if re.search(r"(?:^|[^0-9])\d{2,4}x\d{2,4}(?:[^0-9]|$)", normalized):
        score -= 8
    size_match = re.search(r"-(\d{2,4})-0(?:[^0-9]|$)", raw_url)
    if size_match:
        size = int(size_match.group(1))
        if size <= 320:
            score -= 8
        elif size >= 480:
            score += 5
    return score


def discover_pages(
    session: requests.Session,
    brand_url: str,
    products: list[Product],
    max_pages: int,
    job: JobStatus,
) -> list[PageSnapshot]:
    brand_url = safe_url(brand_url)
    snapshots: list[PageSnapshot] = []
    seen: set[str] = set()
    queue: list[tuple[int, str]] = [(1000, clean_url(brand_url))]
    all_tokens = set().union(*(product_tokens(product) for product in products)) if products else set()

    parsed = urlparse(brand_url)
    search_guesses = [
        f"{parsed.scheme}://{parsed.netloc}/search?q={quote_plus(products[0].name)}" if products else "",
        f"{parsed.scheme}://{parsed.netloc}/buscar?search={quote_plus(products[0].name)}" if products else "",
        f"{parsed.scheme}://{parsed.netloc}/catalogo" if parsed.netloc else "",
        f"{parsed.scheme}://{parsed.netloc}/productos" if parsed.netloc else "",
        f"{parsed.scheme}://{parsed.netloc}/products" if parsed.netloc else "",
    ]
    for guess in search_guesses:
        if guess:
            queue.append((300, guess))

    while queue and len(snapshots) < max_pages:
        queue.sort(key=lambda item: item[0], reverse=True)
        _, url = queue.pop(0)
        if url in seen or not same_domain(brand_url, url):
            continue
        seen.add(url)
        snapshot = snapshot_page(session, url)
        if snapshot is None:
            continue
        snapshots.append(snapshot)

        for href, label in snapshot.links[:300]:
            if href in seen:
                continue
            link_score = score_text_against_tokens(f"{label} {href}", all_tokens)
            catalog_bonus = 4 if re.search(r"producto|product|catalog|shop|item|sku|search|buscar", href, re.I) else 0
            if link_score or catalog_bonus:
                queue.append((link_score + catalog_bonus, href))

    job.add(f"Se analizaron {len(snapshots)} paginas de {brand_url}")
    return snapshots


def ranked_images_for_product(product: Product, pages: list[PageSnapshot]) -> list[ImageCandidate]:
    tokens = product_tokens(product)
    match_tokens = product_match_tokens(product)
    candidates: dict[str, ImageCandidate] = {}

    for page in pages:
        page_path = urlparse(page.url).path
        page_score = score_text_against_tokens(f"{page.title} {page_path}", tokens)
        for image in page.images:
            if any(skip in image.url.lower() for skip in SKIP_IMAGE_HINTS):
                continue
            if not image.url.startswith("http"):
                continue
            match_score = score_text_against_tokens(image.text, match_tokens)
            if match_score < 4:
                continue
            image_score = match_score + min(page_score, 12) + image_quality_score(image.url)
            if product.code and normalize_text(product.code) in normalize_text(image.text):
                image_score += 10
            if not image.url.lower().split("?")[0].endswith(IMAGE_EXTENSIONS):
                image_score -= 2
            if image_score <= 0:
                continue

            existing = candidates.get(image.url)
            if existing is None or image_score > existing.score:
                candidates[image.url] = ImageCandidate(
                    url=image.url,
                    text=image.text,
                    page_url=image.page_url,
                    score=image_score,
                )

    return sorted(candidates.values(), key=lambda item: item.score, reverse=True)


def search_web_images(
    session: requests.Session,
    product: Product,
    serper_api_key: str,
    ai_validation: bool = False,
    limit: int = 20,
) -> tuple[list[ImageCandidate], str]:
    query = product_search_query(product)
    try:
        response = session.post(
            "https://google.serper.dev/images",
            headers={"X-API-KEY": serper_api_key},
            json={"q": query, "gl": "ar", "hl": "es", "num": min(limit, 20)},
            timeout=25,
        )
        response.raise_for_status()
        data = response.json()
    except requests.HTTPError:
        return [], f"Google Images no respondio (HTTP {response.status_code}). Revisa la clave o el saldo de Serper."
    except (requests.RequestException, ValueError):
        return [], "No se pudo consultar Google Images. Revisa la conexion e intenta nuevamente."

    if not isinstance(data, dict):
        return [], "Google Images devolvio una respuesta inesperada."

    results = data.get("images", [])
    candidates = serper_image_candidates(data, product, require_text_match=True)
    if not results:
        return [], f"Google Images no devolvio resultados para: {query}"
    if not candidates:
        return [], f"Google Images devolvio {len(results)} resultados, pero ninguno coincide con el producto: {query}"
    return candidates[:limit], f"Google Images: {len(candidates)} candidatos para: {query}"


def product_search_query(product: Product) -> str:
    brand = re.sub(r"\s+sin\s+stock\b", "", product.brand, flags=re.I).strip()
    identity = product_identity_tokens(product.name, brand)
    return f"{brand} {' '.join(identity)}".strip()[:220]


def google_images_search_url(product: Product) -> str:
    return f"https://www.google.com/search?udm=2&q={quote_plus(product_search_query(product))}"


def serper_image_candidates(
    data: dict, product: Product, require_text_match: bool = True
) -> list[ImageCandidate]:
    tokens = product_match_tokens(product)
    candidates: dict[str, ImageCandidate] = {}

    for result in data.get("images", []):
        if not isinstance(result, dict):
            continue
        image_url = str(result.get("imageUrl") or result.get("image") or "").strip()
        if not image_url.startswith("http"):
            continue
        page_url = str(result.get("link") or result.get("source") or "").strip() or image_url
        text = " ".join(
            [
                str(result.get("title", "")),
                str(result.get("source", "")),
                image_url,
            ]
        )
        match_score = score_text_against_tokens(text, tokens)
        if (require_text_match and match_score < 4) or any(
            skip in image_url.lower() for skip in SKIP_IMAGE_HINTS
        ):
            continue
        score = match_score + image_quality_score(image_url) + 10
        candidates[image_url] = ImageCandidate(
            url=image_url,
            text=text,
            page_url=page_url,
            score=score,
        )

    return sorted(candidates.values(), key=lambda item: item.score, reverse=True)


def extension_from_response(url: str, content_type: str) -> str:
    path_ext = Path(urlparse(url).path).suffix.lower()
    if path_ext in IMAGE_EXTENSIONS:
        return path_ext
    if "png" in content_type:
        return ".png"
    if "webp" in content_type:
        return ".webp"
    if "gif" in content_type:
        return ".gif"
    return ".jpg"


def download_image(session: requests.Session, candidate: ImageCandidate, target: Path) -> Path | None:
    try:
        response = session.get(candidate.url, timeout=20, stream=True)
        response.raise_for_status()
    except requests.RequestException:
        return None

    content_type = response.headers.get("content-type", "")
    if "image" not in content_type and not candidate.url.lower().split("?")[0].endswith(IMAGE_EXTENSIONS):
        return None

    extension = extension_from_response(candidate.url, content_type)
    final_target = target.with_suffix(extension)
    with final_target.open("wb") as file:
        for chunk in response.iter_content(chunk_size=64 * 1024):
            if chunk:
                file.write(chunk)

    try:
        with Image.open(final_target) as image:
            width, height = image.size
    except (OSError, ValueError):
        final_target.unlink(missing_ok=True)
        return None

    # Avoid keeping catalog thumbnails; try the next ranked candidate instead.
    if max(width, height) < 500 or min(width, height) < 150:
        final_target.unlink(missing_ok=True)
        return None

    return final_target if final_target.stat().st_size > 1024 else None


def gemini_mime_type(image_path: Path) -> str:
    return {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".webp": "image/webp",
        ".gif": "image/gif",
        ".bmp": "image/bmp",
    }.get(image_path.suffix.lower(), "image/jpeg")


def validate_image_with_gemini(
    session: requests.Session,
    api_key: str,
    image_path: Path,
    product: Product,
) -> AiVerdict:
    if not api_key:
        return AiVerdict(None, reason="No hay una clave de Gemini configurada.")

    image_data = image_path.read_bytes()
    if len(image_data) > 8 * 1024 * 1024:
        return AiVerdict(None, reason="La imagen supera el limite de analisis de IA.")

    prompt = (
        "Verifica si la imagen muestra exactamente el producto indicado. "
        "Compara marca, nombre, variedad, presentacion y texto visible. "
        "No aceptes otro producto de la misma marca ni una imagen generica. "
        "Responde solo JSON valido con las claves match (boolean), confidence "
        "(entero 0 a 100) y reason (texto corto en espanol).\n\n"
        f"Marca: {product.brand}\n"
        f"Producto: {product.name}\n"
        f"Codigo: {product.code or 'No informado'}"
    )
    payload = {
        "contents": [
            {
                "parts": [
                    {"text": prompt},
                    {
                        "inline_data": {
                            "mime_type": gemini_mime_type(image_path),
                            "data": base64.b64encode(image_data).decode("ascii"),
                        }
                    },
                ]
            }
        ],
        "generationConfig": {"temperature": 0, "responseMimeType": "application/json"},
    }

    for model in ("gemini-3.5-flash-lite", "gemini-2.5-flash-lite"):
        endpoint = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
        try:
            response = session.post(endpoint, params={"key": api_key}, json=payload, timeout=45)
            if response.status_code in (404, 400):
                continue
            if response.status_code == 429:
                return AiVerdict(None, reason="Cuota de Gemini agotada; intenta mas tarde.")
            if response.status_code in (401, 403):
                return AiVerdict(None, reason="Gemini rechazo la clave; revisa GEMINI_API_KEY.")
            response.raise_for_status()
            body = response.json()
            text = body["candidates"][0]["content"]["parts"][0]["text"]
            result = json.loads(text)
            confidence = max(0, min(100, int(result.get("confidence", 0))))
            reason = str(result.get("reason", "Sin detalle.")).strip()[:240]
            accepted = bool(result.get("match", False)) and confidence >= 70
            return AiVerdict(accepted, confidence, reason)
        except (requests.RequestException, KeyError, IndexError, TypeError, ValueError, json.JSONDecodeError):
            return AiVerdict(None, reason="Gemini no pudo validar esta imagen.")

    return AiVerdict(None, reason="El modelo de Gemini no esta disponible para esta clave.")


def download_and_validate_image(
    session: requests.Session,
    candidate: ImageCandidate,
    target: Path,
    product: Product,
    api_key: str,
    ai_validation: bool,
) -> tuple[Path | None, AiVerdict | None]:
    image_path = download_image(session, candidate, target)
    if image_path is None or not ai_validation:
        return image_path, None

    verdict = validate_image_with_gemini(session, api_key, image_path, product)
    if verdict.accepted is False:
        image_path.unlink(missing_ok=True)
        return None, verdict
    return image_path, verdict


def zip_directory(source: Path, target_zip: Path) -> None:
    with zipfile.ZipFile(target_zip, "w", compression=zipfile.ZIP_DEFLATED) as zip_file:
        for file in source.rglob("*"):
            if file.is_file():
                zip_file.write(file, file.relative_to(source))


def run_download_job(
    job_id: str,
    excel_path: Path,
    header_row: int,
    brand_column: str,
    product_column: str,
    code_column: str,
    extra_column: str,
    brand_urls: dict[str, str],
    brand_enabled: dict[str, bool],
    internet_search: bool,
    ai_validation: bool,
    serper_api_key: str,
    max_pages: int,
    max_images: int,
    min_score: int,
) -> None:
    with JOBS_LOCK:
        job = JOBS[job_id]
        job.state = "running"
        job.progress = 1
        job.add("Leyendo Excel")

    try:
        products = parse_excel(excel_path, header_row, brand_column, product_column, code_column, extra_column)
        groups = grouped_brands(products)
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        output_dir = DOWNLOAD_DIR / f"{timestamp}-{job_id[:8]}"
        output_dir.mkdir(parents=True, exist_ok=True)
        session = make_session()
        gemini_api_key = get_config_value("GEMINI_API_KEY")
        if ai_validation and not gemini_api_key:
            ai_validation = False
            job.add("Validacion con IA desactivada: falta GEMINI_API_KEY.")
        elif ai_validation:
            job.add("Validacion con IA activada para las imagenes candidatas.")

        selected_groups = {
            brand: brand_products
            for brand, brand_products in groups.items()
            if brand_enabled.get(brand, False)
        }
        total_products = sum(len(items) for items in selected_groups.values())
        processed_products = 0
        used_image_urls: set[str] = set()
        job.output_dir = output_dir
        job.add(
            f"{len(products)} productos detectados en {len(groups)} marcas; "
            f"se descargaran {total_products} productos de {len(selected_groups)} marcas"
        )

        for brand in groups:
            if brand not in selected_groups:
                job.add(f"Marca no seleccionada: {brand}")

        for brand_index, (brand, brand_products) in enumerate(selected_groups.items(), start=1):
            brand_url = brand_urls.get(brand, "").strip()
            if not brand_url:
                if not internet_search:
                    job.add(f"Marca omitida sin web: {brand}")
                    processed_products += len(brand_products)
                    continue
                job.add(f"Sin web para {brand}; se usara la busqueda en Internet")
                pages = []
            else:
                job.add(f"Analizando {brand} ({brand_index}/{len(selected_groups)})")
                pages = discover_pages(session, brand_url, brand_products, max_pages=max_pages, job=job)
            brand_folder = output_dir / slugify(brand, "marca")
            brand_folder.mkdir(parents=True, exist_ok=True)

            for product in brand_products:
                product_label = f"{product.code} {product.name}".strip()
                product_folder = brand_folder / slugify(product_label, f"fila-{product.row}")
                product_folder.mkdir(parents=True, exist_ok=True)

                ranked = [
                    candidate
                    for candidate in ranked_images_for_product(product, pages)
                    if candidate.score >= min_score
                ][: max_images * 3]
                downloaded = 0

                for candidate_index, candidate in enumerate(ranked, start=1):
                    if downloaded >= max_images:
                        break
                    if candidate.url in used_image_urls:
                        continue
                    target = product_folder / f"{downloaded + 1:02d}-score-{candidate.score}"
                    image_path, verdict = download_and_validate_image(
                        session,
                        candidate,
                        target,
                        product,
                        gemini_api_key,
                        ai_validation,
                    )
                    if verdict and verdict.accepted is False:
                        job.add(f"IA descarto: {brand} - {product_label} ({verdict.reason})")
                    elif verdict and verdict.accepted is None:
                        ai_validation = False
                        job.add(
                            "IA no disponible; se conserva la imagen por coincidencia de texto "
                            f"y se continua sin IA. ({verdict.reason})"
                        )
                    if image_path:
                        downloaded += 1
                        used_image_urls.add(candidate.url)
                        source_note = product_folder / f"{downloaded:02d}-source.txt"
                        ai_note = ""
                        if verdict:
                            state = "aceptada" if verdict.accepted else "sin respuesta"
                            ai_note = f"IA: {state} ({verdict.confidence}%) - {verdict.reason}\n"
                        source_note.write_text(
                            f"Producto: {product_label}\n"
                            f"Marca: {brand}\n"
                            f"Imagen: {candidate.url}\n"
                            f"Pagina: {candidate.page_url}\n"
                            f"Puntaje: {candidate.score}\n"
                            f"{ai_note}",
                            encoding="utf-8",
                        )
                    time.sleep(0.2)

                if internet_search and downloaded < max_images:
                    job.add(f"Buscando imagenes en Google Images: {brand} - {product_label}")
                    web_candidates, search_message = search_web_images(
                        session, product, serper_api_key,
                        ai_validation=ai_validation,
                        limit=max_images * 4,
                    )
                    job.add(search_message)
                    known_urls = {candidate.url for candidate in ranked}
                    for candidate in web_candidates:
                        if downloaded >= max_images:
                            break
                        if candidate.url in known_urls:
                            continue
                        known_urls.add(candidate.url)
                        if candidate.url in used_image_urls:
                            continue
                        target = product_folder / f"{downloaded + 1:02d}-score-{candidate.score}"
                        image_path, verdict = download_and_validate_image(
                            session,
                            candidate,
                            target,
                            product,
                            gemini_api_key,
                            ai_validation,
                        )
                        if verdict and verdict.accepted is False:
                            job.add(f"IA descarto: {brand} - {product_label} ({verdict.reason})")
                        elif verdict and verdict.accepted is None:
                            ai_validation = False
                            job.add(
                                "IA no disponible; se conserva la imagen por coincidencia de texto "
                                f"y se continua sin IA. ({verdict.reason})"
                            )
                        if image_path:
                            downloaded += 1
                            used_image_urls.add(candidate.url)
                            source_note = product_folder / f"{downloaded:02d}-source.txt"
                            ai_note = ""
                            if verdict:
                                state = "aceptada" if verdict.accepted else "sin respuesta"
                                ai_note = f"IA: {state} ({verdict.confidence}%) - {verdict.reason}\n"
                            source_note.write_text(
                                f"Producto: {product_label}\n"
                                f"Marca: {brand}\n"
                                f"Imagen: {candidate.url}\n"
                                f"Pagina: {candidate.page_url}\n"
                                f"Puntaje: {candidate.score}\n"
                                "Origen: Busqueda en Internet\n"
                                f"{ai_note}",
                                encoding="utf-8",
                            )
                        time.sleep(0.2)

                if downloaded == 0:
                    search_url = google_images_search_url(product)
                    (product_folder / "SIN_IMAGEN.txt").write_text(
                        "No se encontro una imagen con puntaje suficiente.\n"
                        "Busqueda manual en Google Images:\n"
                        f"{search_url}\n",
                        encoding="utf-8",
                    )
                    job.missing_searches.append((f"{brand} - {product_label}", search_url))
                    job.add(f"Sin imagen: {brand} - {product_label}")
                else:
                    job.add(f"{downloaded} imagen(es): {brand} - {product_label}")

                processed_products += 1
                job.progress = min(95, int((processed_products / max(total_products, 1)) * 90) + 5)

        zip_path = output_dir.with_suffix(".zip")
        zip_directory(output_dir, zip_path)
        job.zip_path = zip_path
        job.progress = 100
        job.state = "finished"
        job.add("Descarga finalizada")
    except Exception as exc:
        job.state = "failed"
        job.error = str(exc)
        job.add(f"Error: {exc}")


def available_download_runs() -> list[dict[str, int | str]]:
    runs: list[dict[str, int | str]] = []
    for folder in sorted(DOWNLOAD_DIR.iterdir(), key=lambda item: item.name, reverse=True):
        if not folder.is_dir():
            continue
        missing_count = sum(1 for _ in folder.rglob("SIN_IMAGEN.txt"))
        if missing_count:
            runs.append({"name": folder.name, "missing_count": missing_count})
    return runs


def saved_pintofruta_reports() -> list[dict[str, str]]:
    prefixes = (
        "pintofruta-",
        "catalogo-no-detectado-instagram-",
        "productos-inactivados-instagram-",
    )
    reports = [
        path for path in JOB_DIR.glob("*.csv")
        if path.name.startswith(prefixes)
    ]
    return [
        {
            "name": path.name,
            "modified": datetime.fromtimestamp(path.stat().st_mtime).strftime("%d/%m/%Y %H:%M"),
        }
        for path in sorted(reports, key=lambda item: item.stat().st_mtime, reverse=True)[:15]
    ]


def product_from_missing_marker(marker: Path) -> Product | None:
    try:
        relative = marker.relative_to(DOWNLOAD_DIR)
    except ValueError:
        return None
    if len(relative.parts) < 4:
        return None

    brand = readable_slug(marker.parent.parent.name)
    product_name = readable_slug(marker.parent.name)
    try:
        content = marker.read_text(encoding="utf-8")
    except OSError:
        content = ""
    query_match = re.search(r"[?&]q=([^&\s]+)", content)
    if query_match:
        query = unquote_plus(query_match.group(1)).strip()
        brand_prefix = normalize_text(brand)
        if normalize_text(query).startswith(brand_prefix + " "):
            product_name = query[len(brand):].strip()
    return Product(row=0, brand=brand, name=product_name)


def run_retry_missing_images_job(
    job_id: str,
    run_names: list[str],
    serper_api_key: str,
    gemini_api_key: str,
    max_candidates: int,
) -> None:
    with JOBS_LOCK:
        job = JOBS[job_id]
        job.state = "running"
        job.progress = 1
        job.output_dir = DOWNLOAD_DIR
        job.add("Buscando carpetas sin imagen")

    try:
        root = DOWNLOAD_DIR.resolve()
        markers: list[Path] = []
        for run_name in run_names:
            run_folder = (DOWNLOAD_DIR / run_name).resolve()
            if run_folder.parent != root or not run_folder.is_dir():
                continue
            markers.extend(run_folder.rglob("SIN_IMAGEN.txt"))

        if not markers:
            job.progress = 100
            job.state = "finished"
            job.add("No se encontraron productos pendientes en las carpetas seleccionadas")
            return

        session = make_session()
        ai_validation = bool(gemini_api_key)
        if ai_validation:
            job.add(f"Reintentando {len(markers)} productos con Google Images e IA")
        else:
            job.add(f"Reintentando {len(markers)} productos con Google Images; falta la clave de Gemini")

        used_image_urls: set[str] = set()
        for processed, marker in enumerate(markers, start=1):
            product = product_from_missing_marker(marker)
            if product is None:
                job.add(f"No se pudo interpretar la carpeta: {marker.parent.name}")
                continue

            label = f"{product.brand} - {product.name}"
            candidates, search_message = search_web_images(
                session,
                product,
                serper_api_key,
                ai_validation=ai_validation,
                limit=max_candidates,
            )
            job.add(search_message)
            image_path: Path | None = None
            source_candidate: ImageCandidate | None = None
            source_verdict: AiVerdict | None = None

            for candidate in candidates:
                if candidate.url in used_image_urls:
                    continue
                target = marker.parent / "01-reintento"
                image_path, verdict = download_and_validate_image(
                    session,
                    candidate,
                    target,
                    product,
                    gemini_api_key,
                    ai_validation,
                )
                if verdict and verdict.accepted is False:
                    job.add(f"IA descarto: {label} ({verdict.reason})")
                elif verdict and verdict.accepted is None:
                    ai_validation = False
                    job.add(
                        "IA no disponible; se conserva la imagen por coincidencia de texto "
                        f"y se continua sin IA. ({verdict.reason})"
                    )
                if image_path:
                    source_candidate = candidate
                    source_verdict = verdict
                    used_image_urls.add(candidate.url)
                    break
                time.sleep(0.2)

            if image_path and source_candidate:
                source_note = marker.parent / "01-reintento-source.txt"
                ai_note = ""
                if source_verdict:
                    state = "aceptada" if source_verdict.accepted else "sin respuesta"
                    ai_note = f"IA: {state} ({source_verdict.confidence}%) - {source_verdict.reason}\n"
                source_note.write_text(
                    f"Producto: {product.name}\n"
                    f"Marca: {product.brand}\n"
                    f"Imagen: {source_candidate.url}\n"
                    f"Pagina: {source_candidate.page_url}\n"
                    f"Origen: Reintento de Google Images\n"
                    f"{ai_note}",
                    encoding="utf-8",
                )
                found_marker = marker.with_name("IMAGEN_ENCONTRADA.txt")
                found_marker.write_text(
                    "Imagen encontrada durante el reintento.\n"
                    f"Busqueda: {google_images_search_url(product)}\n",
                    encoding="utf-8",
                )
                marker.unlink(missing_ok=True)
                job.add(f"Imagen encontrada: {label}")
            else:
                job.missing_searches.append((label, google_images_search_url(product)))
                job.add(f"Sin imagen despues del reintento: {label}")

            job.progress = min(99, int((processed / len(markers)) * 95) + 4)

        job.progress = 100
        job.state = "finished"
        job.add("Reintento finalizado")
    except Exception as exc:
        job.state = "failed"
        job.error = str(exc)
        job.add(f"Error: {exc}")


COMMERCIAL_IMAGE_NAME = "producto_comercial.png"


def pintofruta_browser_connected() -> bool:
    try:
        response = requests.get(f"{PINFOFRUTA_DEBUG_URL}/json/version", timeout=1)
        return response.ok and bool(response.json().get("webSocketDebuggerUrl"))
    except (requests.RequestException, ValueError):
        return False


def open_pintofruta_browser() -> str:
    if pintofruta_browser_connected():
        return "La sesion de Pinto Fruta ya esta abierta."

    edge_path = next((path for path in PINFOFRUTA_EDGE_CANDIDATES if path.is_file()), None)
    if edge_path is None:
        raise RuntimeError("No se encontro Microsoft Edge para abrir Pinto Fruta")

    BROWSER_PROFILE_DIR.mkdir(parents=True, exist_ok=True)
    subprocess.Popen(
        [
            str(edge_path),
            "--remote-debugging-address=127.0.0.1",
            "--remote-debugging-port=9222",
            f"--user-data-dir={BROWSER_PROFILE_DIR}",
            "--new-window",
            PINFOFRUTA_ADMIN_URL,
        ],
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
    )
    return "Se abrio Edge. Inicia sesion en Pinto Fruta y luego vuelve aqui para comenzar la carga."


def read_database_url(env_path: Path) -> str:
    """Read only DATABASE_URL from a local dotenv file without persisting it."""
    if not env_path.is_file():
        raise RuntimeError("No se encontro el archivo .env indicado")

    for raw_line in env_path.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#"):
            continue
        match = re.match(r"^(?:export\s+)?DATABASE_URL\s*=\s*(.*)$", line, flags=re.I)
        if not match:
            continue
        value = match.group(1).strip()
        if len(value) >= 2 and value[0] == value[-1] and value[0] in {"'", '"'}:
            value = value[1:-1]
        if value:
            return value
    raise RuntimeError("El archivo .env no contiene DATABASE_URL")


def run_pintofruta_empty_brands_job(job_id: str, env_file: str, apply_changes: bool) -> None:
    with JOBS_LOCK:
        job = JOBS[job_id]
        job.state = "running"
        job.progress = 2
        job.add("Conectando a la base de datos de Pinto Fruta")

    try:
        try:
            import psycopg
            from psycopg.rows import dict_row
        except ImportError as exc:
            raise RuntimeError("Falta la dependencia psycopg. Ejecuta: py -m pip install -r requirements.txt") from exc

        database_url = read_database_url(Path(env_file).expanduser())
        with psycopg.connect(database_url, connect_timeout=15, sslmode="require", row_factory=dict_row) as connection:
            with connection.cursor() as cursor:
                cursor.execute("select id, name, active from brands order by name")
                brands = cursor.fetchall()
                cursor.execute(
                    "select distinct brand from products "
                    "where deleted_at is null and lower(coalesce(status, '')) <> 'inactive'"
                )
                active_product_brands = {
                    normalize_text(row["brand"])
                    for row in cursor.fetchall()
                    if normalize_text(row["brand"])
                }

                candidates = [
                    brand
                    for brand in brands
                    if brand["active"] is not False
                    and normalize_text(brand["name"]) not in active_product_brands
                ]
                job.progress = 30
                job.add(
                    f"Encontradas {len(candidates)} marcas activas sin productos activos "
                    f"({len(active_product_brands)} marcas con productos activos)"
                )

                results: list[tuple[str, object, str]] = []
                if apply_changes:
                    total = max(len(candidates), 1)
                    for index, brand in enumerate(candidates, start=1):
                        cursor.execute(
                            "update brands set active = false, featured = false "
                            "where id = %s and active is distinct from false returning id, name",
                            (brand["id"],),
                        )
                        updated = cursor.fetchone()
                        if updated:
                            results.append(("Inactivada", updated["id"], str(updated["name"])))
                            job.add(f"Marca inactivada: {updated['name']}")
                        else:
                            results.append(("Ya estaba inactiva", brand["id"], str(brand["name"])))
                        job.progress = min(95, 30 + int((index / total) * 65))
                    connection.commit()
                else:
                    results = [("Pendiente de inactivar", brand["id"], str(brand["name"])) for brand in candidates]
                    connection.rollback()

        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        report_name = "pintofruta-marcas-inactivadas" if apply_changes else "pintofruta-marcas-sin-productos-activos"
        report_path = JOB_DIR / f"{report_name}-{timestamp}.csv"
        with report_path.open("w", encoding="utf-8-sig", newline="") as file:
            writer = csv.writer(file)
            writer.writerow(["Resultado", "ID", "Marca"])
            writer.writerows(results)

        job.zip_path = report_path
        job.output_dir = JOB_DIR
        job.progress = 100
        job.state = "finished"
        action = "inactivadas" if apply_changes else "encontradas"
        job.add(f"Proceso terminado: {len(results)} marcas {action}. CSV listo para descargar")
    except Exception as exc:
        job.state = "failed"
        job.error = str(exc)
        job.add(f"Error: {exc}")


def instagram_products_with_gemini(image_path: Path, post_text: str, api_key: str) -> list[str]:
    prompt = (
        "Identifica los productos comercializados visibles en esta publicacion de Instagram de una distribuidora. "
        "Usa la imagen y el texto. Devuelve JSON estricto con esta forma: "
        '{"products":["marca producto presentacion"]}. '
        "No inventes productos; si no hay uno reconocible, devuelve una lista vacia.\n\n"
        f"Texto publicado: {post_text[:1800]}"
    )
    payload = {
        "contents": [{"parts": [{"text": prompt}, {"inline_data": {
            "mime_type": "image/jpeg",
            "data": base64.b64encode(image_path.read_bytes()).decode("ascii"),
        }}]}],
        "generationConfig": {"temperature": 0, "responseMimeType": "application/json"},
    }
    for model in ("gemini-3.5-flash-lite", "gemini-2.5-flash-lite"):
        response = requests.post(
            f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent",
            params={"key": api_key}, json=payload, timeout=45,
        )
        if response.status_code in (400, 404):
            continue
        response.raise_for_status()
        answer = response.json()["candidates"][0]["content"]["parts"][0]["text"]
        products = json.loads(answer).get("products", [])
        return [str(product).strip() for product in products if str(product).strip()]
    return []


def instagram_catalog_match(product: str, catalog: list[dict[str, object]]) -> dict[str, object] | None:
    wanted = set(product_search_tokens(product))
    best: tuple[float, dict[str, object]] | None = None
    for item in catalog:
        available = set(product_search_tokens(f"{item['brand']} {item['name']}"))
        if not wanted or not available:
            continue
        score = len(wanted & available) / len(wanted)
        if best is None or score > best[0]:
            best = (score, item)
    return best[1] if best and best[0] >= 0.75 else None


def run_pintofruta_instagram_audit_job(job_id: str, api_key: str) -> None:
    with JOBS_LOCK:
        job = JOBS[job_id]
        job.state = "running"
        job.progress = 1
        job.add("Leyendo publicaciones de Instagram y catalogo de Pinto Fruta")

    try:
        if not pintofruta_browser_connected():
            raise RuntimeError("Abre Pinto Fruta e inicia sesion antes de auditar Instagram")
        from playwright.sync_api import sync_playwright

        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        work_dir = JOB_DIR / f"instagram-audit-{timestamp}"
        work_dir.mkdir(exist_ok=False)
        with sync_playwright() as playwright:
            browser = playwright.chromium.connect_over_cdp(PINFOFRUTA_DEBUG_URL, timeout=15_000)
            context = browser.contexts[0]
            profile = context.new_page()
            profile.goto(PINFOFRUTA_INSTAGRAM_URL, wait_until="domcontentloaded", timeout=30_000)
            profile.wait_for_timeout(1_500)
            urls: set[str] = set()
            stable = 0
            for _ in range(100):
                before = len(urls)
                urls.update(profile.locator("a[href*='/p/'], a[href*='/reel/']").evaluate_all("els => els.map(e => e.href)"))
                profile.mouse.wheel(0, 1800)
                profile.wait_for_timeout(350)
                stable = stable + 1 if len(urls) == before else 0
                if stable >= 8:
                    break
            if not urls:
                raise RuntimeError("No se encontraron publicaciones en el perfil de Instagram")

            admin = context.new_page()
            admin.set_default_timeout(15_000)
            catalog = pintofruta_admin_products(admin)
            if not catalog:
                raise RuntimeError("No se encontraron productos en el administrador")
            job.add(f"Publicaciones: {len(urls)}. Productos activos: {len(catalog)}")

            seen_ids: set[int] = set()
            detected: list[tuple[str, str, str]] = []
            post_page = context.new_page()
            for index, post_url in enumerate(sorted(urls), start=1):
                try:
                    post_page.goto(post_url, wait_until="domcontentloaded", timeout=30_000)
                    post_page.wait_for_timeout(700)
                    post_text = post_page.locator("body").inner_text(timeout=10_000)[:4000]
                    screenshot = work_dir / f"{index:03}.jpg"
                    article = post_page.locator("article")
                    (article if article.count() else post_page.locator("body")).screenshot(
                        path=str(screenshot), type="jpeg", quality=75
                    )
                    for detected_product in instagram_products_with_gemini(screenshot, post_text, api_key):
                        match = instagram_catalog_match(detected_product, catalog)
                        if match:
                            seen_ids.add(int(match["id"]))
                            detected.append((post_url, detected_product, str(match["name"])))
                    job.add(f"Publicacion {index}/{len(urls)} analizada")
                except Exception as exc:
                    job.add(f"Publicacion {index}/{len(urls)} omitida: {str(exc)[:130]}")
                job.progress = min(95, int((index / len(urls)) * 90) + 4)

        report_path = JOB_DIR / f"catalogo-no-detectado-instagram-{timestamp}.csv"
        with report_path.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file)
            writer.writerow(["ID Pinto Fruta", "Marca", "Producto", "Resultado"])
            for item in catalog:
                if int(item["id"]) not in seen_ids:
                    writer.writerow([item["id"], item["brand"], item["name"], "No detectado en publicaciones analizadas"])
        with (work_dir / "productos_detectados.csv").open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file)
            writer.writerow(["Publicacion", "Detectado por IA", "Producto catalogo"])
            writer.writerows(detected)

        job.zip_path = report_path
        job.output_dir = work_dir
        job.progress = 100
        job.state = "finished"
        job.add(f"Auditoria terminada: {len(seen_ids)} detectados, {len(catalog) - len(seen_ids)} sin detectar")
    except Exception as exc:
        job.state = "failed"
        job.error = str(exc)
        job.add(f"Error: {exc}")


def run_pintofruta_deactivate_report_job(job_id: str, report_path: str) -> None:
    with JOBS_LOCK:
        job = JOBS[job_id]
        job.state = "running"
        job.progress = 1
        job.add("Leyendo el listado de productos para inactivar")

    try:
        source = Path(report_path)
        with source.open(encoding="utf-8-sig", newline="") as file:
            targets = list(csv.DictReader(file))
        targets = [target for target in targets if target.get("ID Pinto Fruta", "").strip()]
        if not targets:
            raise RuntimeError("El CSV no tiene una columna 'ID Pinto Fruta' con productos")
        if not pintofruta_browser_connected():
            raise RuntimeError("Abre Pinto Fruta e inicia sesion antes de inactivar productos")
        from playwright.sync_api import sync_playwright

        results: list[list[str]] = []
        with sync_playwright() as playwright:
            browser = playwright.chromium.connect_over_cdp(PINFOFRUTA_DEBUG_URL, timeout=15_000)
            context = browser.contexts[0]
            page = context.new_page()
            page.set_default_timeout(15_000)
            page.goto(PINFOFRUTA_ADMIN_URL, wait_until="domcontentloaded", timeout=30_000)
            page.get_by_role("button", name="Productos").first.click()
            page.wait_for_timeout(1_000)
            cancel = page.get_by_role("button", name="Cancelar")
            if cancel.count() and cancel.last.is_visible():
                cancel.last.click()
            search = page.locator("input[placeholder*='SKU']").first
            total = len(targets)
            for index, target in enumerate(targets, start=1):
                product_id = target["ID Pinto Fruta"].strip()
                sku = target.get("SKU", "").strip()
                product_name = target.get("Producto", "").strip()
                brand = target.get("Marca", "").strip()
                try:
                    search.fill(sku or product_name)
                    page.wait_for_timeout(800)
                    rows = page.locator("tbody tr")
                    target_row = None
                    for row_index in range(rows.count()):
                        row = rows.nth(row_index)
                        row_id = re.sub(r"\D", "", row.locator("td").nth(1).inner_text())
                        if row_id == product_id:
                            target_row = row
                            break
                    if target_row is None:
                        results.append(["No encontrado en catalogo activo", product_id, sku, brand, product_name])
                    else:
                        target_row.get_by_role("button", name="Editar").click()
                        active = page.locator("label").filter(
                            has_text=re.compile(r"^ACTIVO$", re.I)
                        ).first.locator("input[type='checkbox']")
                        active.wait_for(state="visible")
                        if active.is_checked():
                            active.uncheck()
                            page.get_by_role("button", name="Guardar").last.click()
                            active.wait_for(state="hidden", timeout=20_000)
                            results.append(["Inactivado", product_id, sku, brand, product_name])
                            job.add(f"Producto inactivado: {product_name}")
                        else:
                            page.get_by_role("button", name="Cancelar").last.click()
                            results.append(["Ya estaba inactivo", product_id, sku, brand, product_name])
                except Exception as exc:
                    results.append([f"Error: {str(exc)[:180]}", product_id, sku, brand, product_name])
                    job.add(f"No se pudo inactivar {product_name}: {exc}")
                job.progress = min(98, int((index / total) * 94) + 4)

        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        audit_path = JOB_DIR / f"productos-inactivados-instagram-{timestamp}.csv"
        with audit_path.open("w", encoding="utf-8-sig", newline="") as file:
            writer = csv.writer(file)
            writer.writerow(["Resultado", "ID Pinto Fruta", "SKU", "Marca", "Producto"])
            writer.writerows(results)
        job.zip_path = audit_path
        job.output_dir = JOB_DIR
        job.progress = 100
        job.state = "finished"
        job.add(f"Inactivacion terminada: {len(results)} productos procesados")
    except Exception as exc:
        job.state = "failed"
        job.error = str(exc)
        job.add(f"Error: {exc}")


def commercial_images_for_upload(root: Path) -> list[Path]:
    return sorted(path for path in root.rglob(COMMERCIAL_IMAGE_NAME) if path.is_file())


def product_upload_label(image_path: Path, root: Path) -> tuple[str, str]:
    product_name = readable_slug(image_path.parent.name)
    try:
        relative = image_path.parent.relative_to(root)
        brand = readable_slug(relative.parts[-2]) if len(relative.parts) > 1 else ""
    except ValueError:
        brand = ""
    return product_name, brand


def product_search_tokens(value: str) -> list[str]:
    normalized = re.sub(
        r"(?<=[a-z])(?=\d)|(?<=\d)(?=[a-z])",
        " ",
        normalize_text(value),
    )
    ignored = STOPWORDS | {
        "caja",
        "pack",
        "unidad",
        "unidades",
        "frasco",
        "botella",
        "pote",
        "envase",
        "doypack",
        "presentacion",
        "u",
        "g",
        "mg",
        "cc",
        "l",
        "lts",
        "oz",
    }
    return [
        token
        for token in normalized.split()
        if token not in ignored and token not in {"stock", "sin"}
    ]


def product_identity_tokens(product_name: str, brand: str) -> list[str]:
    raw_tokens = re.sub(
        r"(?<=[a-z])(?=\d)|(?<=\d)(?=[a-z])",
        " ",
        normalize_text(product_name),
    ).split()
    brand_tokens = [token for token in product_search_tokens(brand) if token not in {"stock", "sin"}]
    if raw_tokens[: len(brand_tokens)] == brand_tokens:
        raw_tokens = raw_tokens[len(brand_tokens) :]

    presentation_units = {
        "ml", "cc", "gr", "g", "kg", "mg", "lt", "lts", "l",
        "capsula", "capsulas", "u", "un", "unidad", "unidades",
    }
    identity: list[str] = []
    for index, token in enumerate(raw_tokens):
        identity.append(token)
        if token in presentation_units and index > 0 and raw_tokens[index - 1].isdigit():
            break

    return identity or raw_tokens[:8]


def pintofruta_match_tokens(product_name: str, brand: str) -> list[str]:
    identity = " ".join(product_identity_tokens(product_name, brand))
    # La marca se usa para buscar, pero algunos titulos del administrador no
    # la imprimen. Exigirla en la puntuacion puede empatar la variante correcta
    # con otra de sabor distinto que si muestra la marca.
    return product_search_tokens(identity)


def product_search_queries(product_name: str, brand: str) -> list[str]:
    identity = product_identity_tokens(product_name, brand)
    brand_tokens = product_search_tokens(brand)
    queries: list[str] = []
    full_identity = " ".join(identity)
    full_with_brand = " ".join(brand_tokens + identity)
    for query in (full_with_brand, full_identity):
        if query and query not in queries:
            queries.append(query)
    for size in (4, 3, 2, 1):
        query = " ".join(identity[:size])
        if query and query not in queries:
            queries.append(query)
    brand_query = " ".join(brand_tokens)
    if brand_query and brand_query not in queries:
        queries.append(brand_query)
    return queries


def product_card_score(core_tokens: list[str], card_text: str) -> tuple[float, int]:
    available = set(product_search_tokens(card_text))
    if not core_tokens:
        return 0.0, 0

    generic_terms = {
        "barra",
        "bar",
        "chocolate",
        "organico",
        "vegano",
        "keto",
        "proteina",
        "protein",
        "iron",
        "low",
        "carb",
        "aceite",
        "extracto",
        "capsulas",
        "cacao",
        "natural",
        "azucar",
        "sin",
    }
    matched = 0
    matched_weight = 0.0
    total_weight = 0.0
    for token in core_tokens:
        # Los numeros suelen ser la presentacion (60 ml, 24 u, 100 g) y
        # distinguen productos con nombres o sabores muy parecidos.
        if token.isdigit():
            weight = 3.0
        elif token in generic_terms:
            weight = 0.5
        else:
            # Sabores, ingredientes y variedades separan productos de una
            # misma linea aun cuando compartan marca y presentacion.
            weight = 3.0
        total_weight += weight
        if token in available:
            matched += 1
            matched_weight += weight
            continue
        if len(token) >= 4 and any(
            SequenceMatcher(None, token, candidate).ratio() >= 0.85
            for candidate in available
        ):
            matched += 1
            matched_weight += weight * 0.85

    return matched_weight / total_weight, matched


def card_product_name(card_text: str) -> str:
    ignored = {"con imagen", "sin imagen", "cargar imagen", "reemplazar", "borrar"}
    lines = [line.strip() for line in card_text.splitlines() if line.strip()]
    title_lines = [line for line in lines if normalize_text(line) not in ignored]
    return " ".join(title_lines)


def gemini_product_image_part(image_path: Path) -> dict[str, object] | None:
    try:
        with Image.open(image_path) as source:
            image = ImageOps.exif_transpose(source).convert("RGBA")
        image.thumbnail((1200, 1200))
        background = Image.new("RGBA", image.size, "white")
        background.alpha_composite(image)
        payload = BytesIO()
        background.convert("RGB").save(payload, "JPEG", quality=88, optimize=True)
        data = payload.getvalue()
        if not data or len(data) > 4 * 1024 * 1024:
            return None
        return {
            "inline_data": {
                "mime_type": "image/jpeg",
                "data": base64.b64encode(data).decode("ascii"),
            }
        }
    except (OSError, ValueError):
        return None


def choose_product_card_with_gemini(
    api_key: str,
    brand: str,
    product_name: str,
    source_image: Path | None,
    candidates: list[tuple[float, int, object]],
) -> AiCardSelection:
    if not api_key:
        return AiCardSelection(None, reason="IA no configurada")

    options = []
    for index, (score, _, card) in enumerate(candidates[:8]):
        options.append(
            {
                "id": index,
                "nombre": card_product_name(card.inner_text(timeout=3_000)),
                "puntaje_texto": round(score, 3),
            }
        )
    product_identity = " ".join(product_identity_tokens(product_name, brand))
    prompt = (
        "Elige la ficha que representa exactamente el mismo producto. "
        "Compara solo marca, producto, variedad y presentacion. Ignora descripciones "
        "comerciales largas. No inventes equivalencias ni elijas si faltan datos para "
        "distinguir dos fichas duplicadas. Tambien inspecciona la imagen comercial "
        "adjunta: el texto, sabor, variedad, envase o gramaje visibles pueden resolver "
        "una diferencia de nombres. Responde JSON estricto con selected_id "
        "(numero o null), confidence (0 a 100) y reason (texto breve).\n\n"
        f"Marca de origen: {brand}\n"
        f"Producto y presentacion de origen: {product_identity}\n"
        f"Fichas candidatas: {json.dumps(options, ensure_ascii=False)}"
    )
    parts: list[dict[str, object]] = [{"text": prompt}]
    if source_image is not None:
        image_part = gemini_product_image_part(source_image)
        if image_part is not None:
            parts.extend([
                {"text": "Imagen comercial del producto de origen:"},
                image_part,
            ])
    payload = {
        "contents": [{"parts": parts}],
        "generationConfig": {"temperature": 0, "responseMimeType": "application/json"},
    }
    for model in ("gemini-3.5-flash-lite", "gemini-2.5-flash-lite"):
        endpoint = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
        try:
            response = requests.post(endpoint, params={"key": api_key}, json=payload, timeout=30)
            if response.status_code in (400, 404):
                continue
            if response.status_code == 429:
                return AiCardSelection(None, reason="cuota de Gemini agotada")
            if response.status_code in (401, 403):
                return AiCardSelection(None, reason="Gemini rechazo la clave")
            response.raise_for_status()
            text = response.json()["candidates"][0]["content"]["parts"][0]["text"]
            result = json.loads(text)
            selected_id = result.get("selected_id")
            confidence = max(0, min(100, int(result.get("confidence", 0))))
            reason = str(result.get("reason", "Sin detalle.")).strip()[:180]
            if isinstance(selected_id, int) and 0 <= selected_id < len(options) and confidence >= 80:
                return AiCardSelection(selected_id, confidence, reason)
            return AiCardSelection(None, confidence, reason or "la IA no pudo elegir con seguridad")
        except (requests.RequestException, KeyError, IndexError, TypeError, ValueError, json.JSONDecodeError):
            continue
    return AiCardSelection(None, reason="Gemini no estuvo disponible")


def find_matching_product_card(
    page,
    core_tokens: list[str],
    brand: str = "",
    product_name: str = "",
    source_image: Path | None = None,
    gemini_api_key: str = "",
):
    cards = page.locator("article")
    candidates: list[tuple[float, int, object]] = []
    for index in range(cards.count()):
        card = cards.nth(index)
        if card.locator("input[type='file']").count() == 0:
            continue
        score, matched = product_card_score(core_tokens, card.inner_text(timeout=3_000))
        candidates.append((score, matched, card))

    candidates.sort(key=lambda item: (item[0], item[1]), reverse=True)
    if not candidates:
        return None, "la busqueda no devolvio tarjetas"
    best_score, best_hits, best_card = candidates[0]
    second_score = candidates[1][0] if len(candidates) > 1 else 0.0
    is_only_result = len(candidates) == 1
    enough_evidence = best_hits >= min(2, len(core_tokens)) and best_score >= 0.25
    if is_only_result and best_hits >= 1:
        enough_evidence = True
    if not enough_evidence:
        return None, (
            f"mejor coincidencia {best_hits}/{len(core_tokens)} terminos "
            f"entre {len(candidates)} tarjetas"
        )
    if not is_only_result and best_score - second_score < 0.07:
        ai_selection = choose_product_card_with_gemini(
            gemini_api_key,
            brand,
            product_name,
            source_image,
            candidates,
        )
        if ai_selection.selected_index is not None:
            return (
                candidates[ai_selection.selected_index][2],
                f"IA eligio la ficha ({ai_selection.confidence}%): {ai_selection.reason}",
            )
        return None, (
            "hay mas de una coincidencia parecida "
            f"({best_score:.0%} frente a {second_score:.0%}). IA: {ai_selection.reason}"
        )
    return best_card, ""


def manual_candidates_for_product(
    page,
    core_tokens: list[str],
    filter_query: str,
) -> list[ManualCandidate]:
    cards = page.locator("article")
    ranked: list[tuple[float, int, str]] = []
    for index in range(cards.count()):
        card = cards.nth(index)
        if card.locator("input[type='file']").count() == 0:
            continue
        score, _ = product_card_score(core_tokens, card.inner_text(timeout=3_000))
        ranked.append((score, index, card_product_name(card.inner_text(timeout=3_000))))
    ranked.sort(reverse=True)
    if not ranked:
        return []

    minimum_score = max(0.25, ranked[0][0] - 0.18)
    return [
        ManualCandidate(title=title, filter_query=filter_query, article_index=index)
        for score, index, title in ranked
        if score >= minimum_score
    ][:6]


def open_pintofruta_product_images(page) -> None:
    search = page.locator("input[type='search'][placeholder*='SKU']")
    if search.count() and search.first.is_visible():
        return

    button = page.get_by_role("button", name=re.compile(r"imagenes de productos", re.I))
    if button.count() == 0:
        raise RuntimeError("No se encontro el acceso a Imagenes de productos. Verifica que la sesion sea administradora.")
    button.first.click()
    search.first.wait_for(state="visible", timeout=15_000)


def pintofruta_product_search(page):
    open_pintofruta_product_images(page)
    search = page.locator("input[type='search'][placeholder*='SKU']").first
    search.wait_for(state="visible", timeout=15_000)
    return search


def run_pintofruta_upload_job(job_id: str, source_path: str, ai_matching: bool) -> None:
    with JOBS_LOCK:
        job = JOBS[job_id]
        job.state = "running"
        job.progress = 1
        job.add("Verificando imagenes comerciales")

    try:
        source_root = Path(source_path).expanduser().resolve()
        if not source_root.is_dir():
            raise ValueError("La carpeta de imagenes comerciales no existe")
        image_paths = commercial_images_for_upload(source_root)
        if not image_paths:
            raise ValueError(f"No se encontro ningun {COMMERCIAL_IMAGE_NAME} en la carpeta elegida")
        if not pintofruta_browser_connected():
            raise RuntimeError("Abre Pinto Fruta e inicia sesion antes de comenzar la carga")

        gemini_api_key = get_config_value("GEMINI_API_KEY") if ai_matching else ""
        if ai_matching and gemini_api_key:
            job.add("Gemini resolvera solo las coincidencias ambiguas")
        elif ai_matching:
            job.add("Desempate con IA desactivado: falta GEMINI_API_KEY")

        try:
            from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
            from playwright.sync_api import sync_playwright
        except ImportError as exc:
            raise RuntimeError("Falta instalar Playwright. Actualiza las dependencias de la aplicacion.") from exc

        job.output_dir = source_root
        total = len(image_paths)
        with sync_playwright() as playwright:
            browser = playwright.chromium.connect_over_cdp(
                PINFOFRUTA_DEBUG_URL,
                timeout=15_000,
            )
            contexts = browser.contexts
            if not contexts:
                raise RuntimeError("No se encontro la sesion abierta de Edge")
            context = contexts[0]
            page = next((item for item in context.pages if "pintofruta.com" in item.url), None)
            if page is None:
                page = context.new_page()
            page.set_default_timeout(15_000)
            if not page.url.startswith(PINFOFRUTA_ADMIN_URL):
                page.goto(PINFOFRUTA_ADMIN_URL, wait_until="commit", timeout=15_000)
                page.wait_for_timeout(1_000)
            search = pintofruta_product_search(page)
            job.add(f"Listo para cargar {total} imagenes en Pinto Fruta")

            for processed, image_path in enumerate(image_paths, start=1):
                product_name, brand = product_upload_label(image_path, source_root)
                core_tokens = pintofruta_match_tokens(product_name, brand)
                card = None
                reason = "no se pudo formar una busqueda"
                attempted_queries: list[str] = []
                last_query = ""
                for query in product_search_queries(product_name, brand):
                    last_query = query
                    if not search.is_visible():
                        search = pintofruta_product_search(page)
                    search.fill(query)
                    # La grilla aplica el filtro con demora; consultarla antes
                    # devolvia cero tarjetas aunque el producto existiera.
                    page.wait_for_timeout(1_300)
                    card, reason = find_matching_product_card(
                        page,
                        core_tokens,
                        brand=brand,
                        product_name=product_name,
                        source_image=image_path,
                        gemini_api_key=gemini_api_key,
                    )
                    if card is not None:
                        break
                    attempted_queries.append(f"{query!r}: {reason}")
                if card is None:
                    details = "; ".join(attempted_queries[-3:])
                    job.add(f"Omitida {product_name}: {reason}. Consultas: {details}")
                    choices = manual_candidates_for_product(page, core_tokens, last_query)
                    if choices:
                        job.manual_choices.append(
                            ManualChoice(
                                id=uuid.uuid4().hex,
                                product_name=product_name,
                                brand=brand,
                                source_image=image_path,
                                candidates=choices,
                            )
                        )
                        job.add(f"Pendiente de eleccion manual: {product_name}")
                else:
                    if reason:
                        job.add(f"{product_name}: {reason}")
                    file_input = card.locator("input[type='file']").first
                    try:
                        with page.expect_response(
                            lambda response: "/api/admin/upload-image" in response.url
                            and response.request.method == "POST",
                            timeout=20_000,
                        ) as upload_response:
                            file_input.set_input_files(str(image_path))
                        response = upload_response.value
                        if not response.ok:
                            raise RuntimeError(f"el servidor devolvio {response.status}")
                        page.wait_for_timeout(1_000)
                        job.add(f"Imagen cargada o reemplazada: {product_name}")
                    except PlaywrightTimeoutError:
                        job.add(f"No se confirmo la carga de {product_name}; se deja para revision manual")
                    except Exception as exc:
                        job.add(f"No se pudo cargar {product_name}: {exc}")
                job.progress = min(99, int((processed / total) * 95) + 4)

        job.progress = 100
        job.state = "finished"
        job.add("Carga en Pinto Fruta finalizada")
    except Exception as exc:
        job.state = "failed"
        job.error = str(exc)
        job.add(f"Error: {exc}")


def run_manual_pintofruta_upload_job(
    job_id: str,
    selections: list[tuple[ManualChoice, ManualCandidate]],
) -> None:
    with JOBS_LOCK:
        job = JOBS[job_id]
        job.state = "running"
        job.progress = 1
        job.add("Preparando las elecciones manuales")

    try:
        if not pintofruta_browser_connected():
            raise RuntimeError("Abre Pinto Fruta e inicia sesion antes de comenzar la carga")
        from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
        from playwright.sync_api import sync_playwright

        with sync_playwright() as playwright:
            browser = playwright.chromium.connect_over_cdp(PINFOFRUTA_DEBUG_URL, timeout=15_000)
            context = browser.contexts[0]
            page = next((item for item in context.pages if "pintofruta.com" in item.url), None)
            if page is None:
                raise RuntimeError("No se encontro la sesion de Pinto Fruta")
            page.set_default_timeout(15_000)
            search = pintofruta_product_search(page)
            total = len(selections)

            for processed, (choice, candidate) in enumerate(selections, start=1):
                if not choice.source_image.is_file():
                    job.add(f"No se encontro la imagen de origen: {choice.product_name}")
                    continue
                if not search.is_visible():
                    search = pintofruta_product_search(page)
                search.fill(candidate.filter_query)
                page.wait_for_timeout(1_300)
                cards = page.locator("article")
                selected_card = None
                if candidate.article_index < cards.count():
                    candidate_card = cards.nth(candidate.article_index)
                    if normalize_text(card_product_name(candidate_card.inner_text())) == normalize_text(candidate.title):
                        selected_card = candidate_card
                if selected_card is None:
                    matching_cards = [
                        cards.nth(index)
                        for index in range(cards.count())
                        if normalize_text(card_product_name(cards.nth(index).inner_text())) == normalize_text(candidate.title)
                    ]
                    if len(matching_cards) == 1:
                        selected_card = matching_cards[0]
                if selected_card is None or selected_card.locator("input[type='file']").count() == 0:
                    job.add(f"La ficha elegida ya no coincide: {choice.product_name}")
                    continue

                try:
                    with page.expect_response(
                        lambda response: "/api/admin/upload-image" in response.url
                        and response.request.method == "POST",
                        timeout=20_000,
                    ) as upload_response:
                        selected_card.locator("input[type='file']").first.set_input_files(str(choice.source_image))
                    response = upload_response.value
                    if not response.ok:
                        raise RuntimeError(f"el servidor devolvio {response.status}")
                    page.wait_for_timeout(1_000)
                    job.add(f"Eleccion cargada: {choice.product_name}")
                except PlaywrightTimeoutError:
                    job.add(f"No se confirmo la carga de {choice.product_name}")
                except Exception as exc:
                    job.add(f"No se pudo cargar {choice.product_name}: {exc}")
                job.progress = min(99, int((processed / total) * 95) + 4)

        job.progress = 100
        job.state = "finished"
        job.add("Elecciones manuales finalizadas")
    except Exception as exc:
        job.state = "failed"
        job.error = str(exc)
        job.add(f"Error: {exc}")


def run_pintofruta_missing_report_job(job_id: str) -> None:
    with JOBS_LOCK:
        job = JOBS[job_id]
        job.state = "running"
        job.progress = 1
        job.add("Leyendo el catalogo de Pinto Fruta")

    try:
        if not pintofruta_browser_connected():
            raise RuntimeError("Abre Pinto Fruta e inicia sesion antes de generar el listado")
        from playwright.sync_api import sync_playwright

        with sync_playwright() as playwright:
            browser = playwright.chromium.connect_over_cdp(PINFOFRUTA_DEBUG_URL, timeout=15_000)
            context = browser.contexts[0]
            page = next((item for item in context.pages if "pintofruta.com" in item.url), None)
            if page is None:
                raise RuntimeError("No se encontro la sesion de Pinto Fruta")
            page.set_default_timeout(15_000)
            search = pintofruta_product_search(page)
            search.fill("")
            page.wait_for_timeout(1_300)
            cards = page.locator("article")
            missing_products: list[str] = []
            for index in range(cards.count()):
                card_text = cards.nth(index).inner_text(timeout=3_000)
                if "sin imagen" in normalize_text(card_text):
                    missing_products.append(card_product_name(card_text))
                job.progress = min(95, int(((index + 1) / max(cards.count(), 1)) * 90) + 4)

        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        report_path = JOB_DIR / f"pintofruta-sin-imagen-{timestamp}.csv"
        with report_path.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file)
            writer.writerow(["Producto", "Busqueda Google Images"])
            for product_name in sorted(set(missing_products), key=normalize_text):
                writer.writerow([
                    product_name,
                    f"https://www.google.com/search?tbm=isch&q={quote_plus(product_name)}",
                ])
        job.zip_path = report_path
        job.progress = 100
        job.state = "finished"
        job.add(f"Listado creado: {len(missing_products)} productos sin imagen")
    except Exception as exc:
        job.state = "failed"
        job.error = str(exc)
        job.add(f"Error: {exc}")


def pintofruta_duplicate_key(product_name: str) -> str:
    identity = " ".join(product_identity_tokens(product_name, ""))
    return identity or normalize_text(product_name)


def pintofruta_admin_products(page) -> list[dict[str, object]]:
    page.goto(PINFOFRUTA_ADMIN_URL, wait_until="commit", timeout=15_000)
    page.wait_for_timeout(1_000)
    page.get_by_role("button", name="Productos").first.click()
    page.wait_for_timeout(600)
    products: list[dict[str, object]] = []
    while True:
        rows = page.locator("tbody tr").evaluate_all(
            "els => els.map(e => [...e.querySelectorAll('td')].map(td => td.innerText.trim()))"
        )
        for row in rows:
            if len(row) < 5:
                continue
            product_id = re.sub(r"\D", "", row[1])
            if product_id:
                products.append({
                    "id": int(product_id),
                    "sku": row[2],
                    "name": row[3],
                    "brand": row[4],
                })
        next_button = page.get_by_role("button", name="Siguiente")
        if next_button.is_disabled():
            break
        previous_id = rows[0][1] if rows else ""
        next_button.click()
        page.wait_for_function(
            "old => document.querySelector('tbody tr td:nth-child(2)')?.innerText.trim() !== old",
            arg=previous_id,
            timeout=15_000,
        )
    return products


def run_pintofruta_duplicate_cleanup_job(job_id: str) -> None:
    with JOBS_LOCK:
        job = JOBS[job_id]
        job.state = "running"
        job.progress = 1
        job.add("Analizando duplicados y sus imagenes")

    try:
        if not pintofruta_browser_connected():
            raise RuntimeError("Abre Pinto Fruta e inicia sesion antes de eliminar duplicados")
        from playwright.sync_api import sync_playwright

        with sync_playwright() as playwright:
            browser = playwright.chromium.connect_over_cdp(PINFOFRUTA_DEBUG_URL, timeout=15_000)
            context = browser.contexts[0]
            catalog_page = context.new_page()
            catalog_page.set_default_timeout(15_000)
            catalog_page.goto(PINFOFRUTA_ADMIN_URL, wait_until="domcontentloaded", timeout=30_000)
            catalog_search = pintofruta_product_search(catalog_page)
            catalog_search.fill("")
            catalog_page.wait_for_timeout(1_500)
            cards = catalog_page.locator("article").evaluate_all(
                "els => els.map(e => ({text: e.innerText, src: e.querySelector('img')?.getAttribute('src') || ''}))"
            )

            admin_page = context.new_page()
            admin_page.set_default_timeout(15_000)
            products = pintofruta_admin_products(admin_page)
            if not products:
                raise RuntimeError("No se encontraron productos en el administrador")
            known_ids = {int(product["id"]) for product in products}
            image_ids: set[int] = set()
            for card in cards:
                if "con imagen" not in normalize_text(card["text"]):
                    continue
                image_ids.update(
                    int(value)
                    for value in re.findall(r"-(\d+)(?:-|\.)", card["src"])
                    if int(value) in known_ids
                )

            grouped: dict[tuple[str, str], list[dict[str, object]]] = {}
            for product in products:
                key = duplicate_product_key(Product(0, str(product["brand"]), str(product["name"])))
                grouped.setdefault(key, []).append(product)
            targets: list[dict[str, object]] = []
            for group in grouped.values():
                if len(group) < 2:
                    continue
                keep = max(group, key=lambda product: (int(product["id"]) in image_ids, int(product["id"])))
                targets.extend(product for product in group if product["id"] != keep["id"])

            job.add(f"Se eliminaran {len(targets)} duplicados. Se conserva imagen; en empate, el ID mas reciente.")
            report_path = JOB_DIR / f"pintofruta-duplicados-eliminados-{datetime.now().strftime('%Y%m%d-%H%M%S')}.csv"
            deleted: list[dict[str, object]] = []
            errors: list[tuple[dict[str, object], str]] = []
            admin_page.on("dialog", lambda dialog: dialog.accept())
            search = admin_page.locator("input[placeholder*='SKU']").first
            total = len(targets)
            for index, product in enumerate(targets, start=1):
                try:
                    search.fill(str(product["sku"]))
                    admin_page.wait_for_timeout(400)
                    rows = admin_page.locator("tbody tr")
                    target_row = None
                    for row_index in range(rows.count()):
                        row = rows.nth(row_index)
                        row_id = re.sub(r"\D", "", row.locator("td").nth(1).inner_text())
                        if row_id == str(product["id"]):
                            target_row = row
                            break
                    if target_row is None:
                        raise RuntimeError("no se encontro la fila por ID")
                    target_row.get_by_role("button", name="Borrar").click()
                    admin_page.wait_for_timeout(650)
                    deleted.append(product)
                    job.add(f"Eliminado ID {product['id']}: {product['name']}")
                except Exception as exc:
                    errors.append((product, str(exc)))
                    job.add(f"No se pudo eliminar ID {product['id']}: {exc}")
                job.progress = min(99, int((index / max(total, 1)) * 92) + 6)

        with report_path.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file)
            writer.writerow(["Resultado", "ID", "SKU", "Marca", "Producto", "Tenia imagen"])
            for product in deleted:
                writer.writerow(["Eliminado", product["id"], product["sku"], product["brand"], product["name"], int(product["id"]) in image_ids])
            for product, error in errors:
                writer.writerow([f"Error: {error}", product["id"], product["sku"], product["brand"], product["name"], int(product["id"]) in image_ids])
        job.zip_path = report_path
        job.progress = 100
        job.state = "finished"
        job.add(f"Limpieza finalizada: {len(deleted)} eliminados, {len(errors)} con error")
    except Exception as exc:
        job.state = "failed"
        job.error = str(exc)
        job.add(f"Error: {exc}")


def run_pintofruta_duplicate_report_job(job_id: str) -> None:
    with JOBS_LOCK:
        job = JOBS[job_id]
        job.state = "running"
        job.progress = 1
        job.add("Leyendo el catalogo de Pinto Fruta")

    try:
        if not pintofruta_browser_connected():
            raise RuntimeError("Abre Pinto Fruta e inicia sesion antes de generar el listado")
        from playwright.sync_api import sync_playwright

        with sync_playwright() as playwright:
            browser = playwright.chromium.connect_over_cdp(PINFOFRUTA_DEBUG_URL, timeout=15_000)
            context = browser.contexts[0]
            page = next((item for item in context.pages if "pintofruta.com" in item.url), None)
            if page is None:
                raise RuntimeError("No se encontro la sesion de Pinto Fruta")
            page.set_default_timeout(15_000)
            search = pintofruta_product_search(page)
            search.fill("")
            page.wait_for_timeout(1_500)
            cards = page.locator("article")
            product_names: list[str] = []
            total = cards.count()
            if not total:
                raise RuntimeError("No se encontraron productos en el catalogo visual")
            for index in range(total):
                name = card_product_name(cards.nth(index).inner_text(timeout=3_000))
                if name:
                    product_names.append(name)
                job.progress = min(90, int(((index + 1) / total) * 80) + 5)

        grouped: dict[str, list[str]] = {}
        for name in product_names:
            grouped.setdefault(pintofruta_duplicate_key(name), []).append(name)
        duplicates = [names for names in grouped.values() if len(names) > 1]
        duplicates.sort(key=lambda names: (pintofruta_duplicate_key(names[0]), normalize_text(names[0])))

        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        report_path = JOB_DIR / f"pintofruta-productos-duplicados-{timestamp}.csv"
        with report_path.open("w", newline="", encoding="utf-8-sig") as file:
            writer = csv.writer(file)
            writer.writerow([
                "Tipo de coincidencia", "Clave de producto y presentacion",
                "Producto en Pinto Fruta", "Cantidad en el catalogo",
            ])
            for names in duplicates:
                normalized_names = {normalize_text(name) for name in names}
                match_type = "Exacta" if len(normalized_names) == 1 else "Mismo producto y presentacion"
                key = pintofruta_duplicate_key(names[0])
                for name in sorted(names, key=normalize_text):
                    writer.writerow([match_type, key, name, len(names)])

        job.zip_path = report_path
        job.progress = 100
        job.state = "finished"
        job.add(f"Listado creado: {len(duplicates)} grupos duplicados entre {len(product_names)} productos")
    except Exception as exc:
        job.state = "failed"
        job.error = str(exc)
        job.add(f"Error: {exc}")


def products_from_listing(path: Path) -> list[Product]:
    rows: list[dict[str, str]] = []
    first_data_row = 2
    if path.suffix.lower() == ".csv":
        content = path.read_text(encoding="utf-8-sig")
        reader = csv.DictReader(content.splitlines())
        rows = [{str(key or ""): str(value or "") for key, value in row.items()} for row in reader]
    elif path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        worksheet = workbook.active
        values = list(worksheet.iter_rows(values_only=True))
        if not values:
            return []
        layout = read_excel_layout(path)
        header_index = layout.header_row - 1
        first_data_row = layout.header_row + 1
        headers = [str(value or "").strip() for value in values[header_index]]
        rows = [
            {headers[index]: str(value or "").strip() for index, value in enumerate(row) if index < len(headers)}
            for row in values[header_index + 1 :]
        ]
    else:
        raise ValueError("El listado debe ser un archivo CSV o XLSX")

    products: list[Product] = []
    for row_number, row in enumerate(rows, start=first_data_row):
        normalized = {normalize_text(key): value.strip() for key, value in row.items()}
        name = next((value for key, value in normalized.items() if value and any(
            term in key for term in ("producto", "nombre", "detalle", "descripcion", "variedad", "product")
        )), "")
        if not name:
            name = next((value for value in normalized.values() if value), "")
        brand = next((value for key, value in normalized.items() if value and (key == "marca" or key == "brand" or "marca" in key)), "")
        if name:
            products.append(Product(row=row_number, brand=brand, name=name))
    return products


def duplicate_product_key(product: Product) -> tuple[str, str]:
    brand = re.sub(r"\bsin\s+stock\b", "", normalize_text(product.brand)).strip()
    identity = " ".join(product_identity_tokens(product.name, product.brand))
    return brand, identity or normalize_text(product.name)


def run_duplicate_listing_job(job_id: str, listing_path: str) -> None:
    with JOBS_LOCK:
        job = JOBS[job_id]
        job.state = "running"
        job.progress = 10
        job.add("Leyendo el listado de productos")

    try:
        source_path = Path(listing_path).resolve()
        if not source_path.is_file():
            raise ValueError("No se encontro el archivo de origen")
        products = products_from_listing(source_path)
        if not products:
            raise ValueError("El archivo no contiene productos")

        grouped: dict[tuple[str, str], list[Product]] = {}
        for product in products:
            grouped.setdefault(duplicate_product_key(product), []).append(product)
        duplicates = [group for group in grouped.values() if len(group) > 1]
        duplicates.sort(key=lambda group: duplicate_product_key(group[0]))
        job.progress = 65
        job.add(f"Se revisaron {len(products)} productos")

        report_path = JOB_DIR / f"productos-duplicados-{datetime.now().strftime('%Y%m%d-%H%M%S')}.csv"
        with report_path.open("w", encoding="utf-8-sig", newline="") as report:
            writer = csv.writer(report)
            writer.writerow(["Clave de coincidencia", "Marca", "Producto", "Fila de origen", "Repeticiones"])
            for group in duplicates:
                key = " - ".join(part for part in duplicate_product_key(group[0]) if part)
                for product in group:
                    writer.writerow([key, product.brand, product.name, product.row, len(group)])

        job.zip_path = report_path
        job.progress = 100
        job.state = "finished"
        job.add(f"Reporte listo: {len(duplicates)} grupos duplicados")
    except Exception as exc:
        job.state = "failed"
        job.error = str(exc)
        job.add(f"Error: {exc}")


def run_listing_image_preparation_job(
    job_id: str,
    listing_path: str,
    destination_path: str,
    max_candidates: int,
) -> None:
    with JOBS_LOCK:
        job = JOBS[job_id]
        job.state = "running"
        job.progress = 1
        job.add("Leyendo el listado de productos")

    try:
        source_path = Path(listing_path).resolve()
        destination_root = Path(destination_path).expanduser().resolve()
        if not source_path.is_file():
            raise ValueError("No se encontro el archivo de origen")
        if not destination_root.is_dir():
            raise ValueError("La carpeta de descarga no existe")
        products = products_from_listing(source_path)
        if not products:
            raise ValueError("El archivo no contiene productos para buscar")
        serper_api_key = get_config_value("SERPER_API_KEY")
        if not serper_api_key:
            raise ValueError("Falta configurar SERPER_API_KEY para buscar en Google Images")
        gemini_api_key = get_config_value("GEMINI_API_KEY")
        output_root = destination_root / (
            f"{slugify(source_path.stem, 'productos')}-preparadas-"
            f"{datetime.now().strftime('%Y%m%d-%H%M%S')}"
        )
        output_root.mkdir(parents=True, exist_ok=False)
        job.output_dir = output_root
        session = make_session()
        try:
            from rembg import new_session
        except ImportError as exc:
            raise RuntimeError("Falta instalar rembg para preparar las imagenes") from exc
        background_session = new_session("u2net")
        job.add(f"Buscando y preparando {len(products)} productos")

        for processed, product in enumerate(products, start=1):
            label = product_search_query(product)
            folder = output_root / slugify(product.name, f"producto-{processed}")
            folder.mkdir(parents=True, exist_ok=True)
            candidates, message = search_web_images(
                session,
                product,
                serper_api_key,
                ai_validation=bool(gemini_api_key),
                limit=max_candidates,
            )
            job.add(message)
            downloaded: Path | None = None
            for candidate in candidates:
                raw_target = folder / "01-origen.jpg"
                image_path, verdict = download_and_validate_image(
                    session,
                    candidate,
                    raw_target,
                    product,
                    gemini_api_key,
                    ai_validation=bool(gemini_api_key),
                )
                if image_path:
                    downloaded = image_path
                    break
                if verdict and verdict.accepted is False:
                    job.add(f"IA descarto: {label} ({verdict.reason})")

            if downloaded is None:
                (folder / "SIN_IMAGEN.txt").write_text(
                    f"Producto: {product.name}\nBusqueda: {google_images_search_url(product)}\n",
                    encoding="utf-8",
                )
                job.add(f"Sin imagen: {label}")
            else:
                try:
                    target = folder / COMMERCIAL_IMAGE_NAME
                    create_commercial_image(downloaded, target, background_session)
                    job.add(f"Lista para subir: {label}")
                except (OSError, ValueError, RuntimeError) as exc:
                    job.add(f"No se pudo preparar {label}: {exc}")
            job.progress = min(99, int((processed / len(products)) * 95) + 4)

        job.progress = 100
        job.state = "finished"
        job.add("Busqueda y preparacion finalizadas")
    except Exception as exc:
        job.state = "failed"
        job.error = str(exc)
        job.add(f"Error: {exc}")


def commercial_source_files(root: Path) -> dict[Path, list[Path]]:
    grouped: dict[Path, list[Path]] = {}
    for image_path in root.rglob("*"):
        if not image_path.is_file() or image_path.suffix.lower() not in IMAGE_EXTENSIONS:
            continue
        if image_path.name.lower() == COMMERCIAL_IMAGE_NAME:
            continue
        grouped.setdefault(image_path.parent, []).append(image_path)
    return grouped


def best_source_image(files: list[Path]) -> Path | None:
    choices: list[tuple[int, int, Path]] = []
    for image_path in files:
        try:
            with Image.open(image_path) as image:
                width, height = image.size
            choices.append((width * height, image_path.stat().st_size, image_path))
        except (OSError, ValueError):
            continue
    return max(choices, default=(0, 0, None))[2]


def create_commercial_image(source: Path, target: Path, background_session: object) -> tuple[int, int]:
    from rembg import remove

    with Image.open(source) as image:
        original = ImageOps.exif_transpose(image).convert("RGBA")
    result = remove(original, session=background_session)
    if isinstance(result, bytes):
        with Image.open(BytesIO(result)) as output:
            result = output.convert("RGBA")
    else:
        result = result.convert("RGBA")

    bounds = result.getchannel("A").getbbox()
    if bounds is None:
        raise ValueError("No se pudo separar el producto del fondo")
    product = result.crop(bounds)
    side = max(product.width, product.height)
    padding = max(24, round(side * 0.08))
    canvas = Image.new("RGBA", (side + padding * 2, side + padding * 2), (255, 255, 255, 0))
    offset = ((canvas.width - product.width) // 2, (canvas.height - product.height) // 2)
    canvas.alpha_composite(product, offset)
    canvas.save(target, "PNG", optimize=True)
    return canvas.size


def run_commercial_images_job(job_id: str, source_path: str, destination_path: str) -> None:
    with JOBS_LOCK:
        job = JOBS[job_id]
        job.state = "running"
        job.progress = 1
        job.add("Buscando imagenes de producto")

    try:
        source_root = Path(source_path).expanduser().resolve()
        destination_root = Path(destination_path).expanduser().resolve()
        if not source_root.is_dir():
            raise ValueError("La carpeta de origen no existe")
        if not destination_root.is_dir():
            raise ValueError("La carpeta de destino no existe")
        if source_root == destination_root:
            raise ValueError("El origen y el destino deben ser carpetas distintas")
        try:
            destination_root.relative_to(source_root)
        except ValueError:
            pass
        else:
            raise ValueError("La carpeta de destino no puede estar dentro del origen")

        output_root = destination_root / (
            f"{slugify(source_root.name, 'productos')}-comerciales-"
            f"{datetime.now().strftime('%Y%m%d-%H%M%S')}"
        )
        output_root.mkdir(parents=True, exist_ok=False)
        job.output_dir = output_root
        sources_by_folder = commercial_source_files(source_root)
        if not sources_by_folder:
            job.progress = 100
            job.state = "finished"
            job.add("No se encontraron imagenes para procesar")
            return

        try:
            from rembg import new_session
        except ImportError as exc:
            raise RuntimeError("Falta instalar rembg. Ejecuta la instalacion de dependencias.") from exc

        job.add(f"Preparando fondo transparente para {len(sources_by_folder)} productos")
        background_session = new_session("u2net")
        total = len(sources_by_folder)
        for processed, (folder, files) in enumerate(sorted(sources_by_folder.items()), start=1):
            source = best_source_image(files)
            label = readable_slug(folder.name)
            if source is None:
                job.add(f"No se pudo abrir una imagen: {label}")
            else:
                try:
                    relative_folder = folder.relative_to(source_root)
                    target_folder = output_root / relative_folder
                    target_folder.mkdir(parents=True, exist_ok=True)
                    target = target_folder / COMMERCIAL_IMAGE_NAME
                    width, height = create_commercial_image(source, target, background_session)
                    job.add(f"Imagen comercial creada: {label} ({width} x {height})")
                except (OSError, ValueError, RuntimeError) as exc:
                    job.add(f"No se pudo procesar {label}: {exc}")
            job.progress = min(99, int((processed / total) * 95) + 4)

        job.progress = 100
        job.state = "finished"
        job.add("Imagenes comerciales finalizadas")
    except Exception as exc:
        job.state = "failed"
        job.error = str(exc)
        job.add(f"Error: {exc}")


BASE_TEMPLATE = """
<!doctype html>
<html lang="es">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>IMG Downloader</title>
  <style>
    :root {
      color-scheme: light;
      --bg: #f6f7f9;
      --panel: #ffffff;
      --ink: #1f2933;
      --muted: #64748b;
      --line: #d9e0e8;
      --accent: #0f766e;
      --accent-dark: #115e59;
      --warn: #a16207;
      --bad: #b91c1c;
    }
    * { box-sizing: border-box; }
    body {
      margin: 0;
      font-family: "Segoe UI", system-ui, -apple-system, BlinkMacSystemFont, sans-serif;
      background: var(--bg);
      color: var(--ink);
    }
    main {
      width: min(1180px, calc(100vw - 32px));
      margin: 32px auto;
    }
    header {
      display: flex;
      align-items: end;
      justify-content: space-between;
      gap: 24px;
      margin-bottom: 24px;
    }
    h1, h2, h3 { margin: 0; letter-spacing: 0; }
    h1 { font-size: clamp(28px, 4vw, 44px); }
    h2 { font-size: 22px; margin-bottom: 16px; }
    h3 { font-size: 16px; }
    p { color: var(--muted); line-height: 1.5; }
    .panel {
      background: var(--panel);
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 22px;
      box-shadow: 0 16px 40px rgba(31, 41, 51, 0.07);
    }
    .grid { display: grid; grid-template-columns: repeat(12, 1fr); gap: 16px; }
    .span-3 { grid-column: span 3; }
    .span-4 { grid-column: span 4; }
    .span-6 { grid-column: span 6; }
    .span-12 { grid-column: span 12; }
    label { display: block; font-weight: 650; margin-bottom: 8px; }
    input, select {
      width: 100%;
      min-height: 42px;
      border: 1px solid var(--line);
      border-radius: 6px;
      padding: 9px 11px;
      font: inherit;
      background: white;
      color: var(--ink);
    }
    input[type="file"] { padding: 8px; }
    input[type="checkbox"].brand-check {
      width: 20px;
      min-height: 20px;
      margin: 0;
      accent-color: var(--accent);
      cursor: pointer;
    }
    input[type="checkbox"].retry-run-check {
      width: 20px;
      min-height: 20px;
      margin: 0;
      accent-color: var(--accent);
      cursor: pointer;
    }
    button, .button {
      display: inline-flex;
      align-items: center;
      justify-content: center;
      min-height: 42px;
      border: 0;
      border-radius: 6px;
      padding: 10px 15px;
      font: inherit;
      font-weight: 700;
      background: var(--accent);
      color: white;
      cursor: pointer;
      text-decoration: none;
    }
    button:hover, .button:hover { background: var(--accent-dark); }
    .secondary { background: #334155; }
    .secondary:hover { background: #1f2937; }
    .hint { font-size: 13px; margin: 6px 0 0; }
    .brand-table {
      width: 100%;
      border-collapse: collapse;
      margin-top: 8px;
    }
    .brand-table th, .brand-table td {
      border-bottom: 1px solid var(--line);
      padding: 10px 8px;
      text-align: left;
      vertical-align: middle;
    }
    .brand-table th { font-size: 13px; color: var(--muted); }
    .brand-table th:first-child, .brand-table td:first-child { width: 92px; text-align: center; }
    .brand-table td:nth-child(2) { width: 30%; font-weight: 650; }
    .brand-table td:nth-child(3) { width: 90px; color: var(--muted); }
    .table-tools { display: flex; justify-content: flex-end; margin: 14px 0 8px; }
    .table-tools button { min-height: 36px; padding: 8px 12px; font-size: 14px; }
    .actions { display: flex; gap: 10px; flex-wrap: wrap; align-items: center; margin-top: 18px; }
    .manual-searches {
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(min(100%, 320px), 1fr));
      gap: 8px;
      margin-top: 10px;
    }
    .manual-search {
      justify-content: flex-start;
      min-width: 0;
      text-align: left;
      white-space: normal;
      overflow-wrap: anywhere;
    }
    .folder-picker {
      display: grid;
      grid-template-columns: minmax(0, 1fr) auto;
      gap: 10px;
    }
    .progress {
      height: 16px;
      background: #e5e7eb;
      border-radius: 999px;
      overflow: hidden;
      border: 1px solid var(--line);
    }
    .bar {
      height: 100%;
      width: var(--value);
      background: var(--accent);
      transition: width 0.3s ease;
    }
    pre {
      max-height: 460px;
      overflow: auto;
      padding: 14px;
      border-radius: 6px;
      background: #101820;
      color: #d7e3ea;
      font-size: 13px;
      line-height: 1.45;
      white-space: pre-wrap;
    }
    .status {
      display: inline-flex;
      align-items: center;
      min-height: 30px;
      padding: 4px 9px;
      border-radius: 999px;
      background: #dff7ef;
      color: #064e3b;
      font-weight: 700;
      font-size: 13px;
    }
    .status.failed { background: #fee2e2; color: var(--bad); }
    .status.running { background: #e0f2fe; color: #075985; }
    @media (max-width: 760px) {
      main { width: min(100vw - 20px, 1180px); margin: 18px auto; }
      header { display: block; }
      .span-3, .span-4, .span-6 { grid-column: span 12; }
      .panel { padding: 16px; }
      .brand-table td:nth-child(2) { width: auto; }
      .folder-picker { grid-template-columns: 1fr; }
    }
  </style>
</head>
<body>
<main>
  <header>
    <div>
      <h1>IMG Downloader</h1>
      <p>Importa productos desde Excel, agrupa por marca y descarga imagenes desde webs de fabricantes.</p>
    </div>
    <a class="button secondary" href="{{ url_for('index') }}">Nuevo proceso</a>
  </header>
  {{ content|safe }}
</main>
</body>
</html>
"""


def page(content: str, **context):
    return render_template_string(BASE_TEMPLATE, content=render_template_string(content, **context))


@app.get("/")
def index():
    return page(
        """
        <section class="panel">
          <h2>1. Subir Excel</h2>
          <form method="post" action="{{ url_for('upload') }}" enctype="multipart/form-data">
            <div class="grid">
              <div class="span-12">
                <label for="excel">Archivo .xlsx</label>
                <input id="excel" name="excel" type="file" accept=".xlsx" required>
                <p class="hint">Puede tener titulos o avisos arriba; la app busca automaticamente la fila de encabezados.</p>
              </div>
            </div>
            <div class="actions">
              <button type="submit">Continuar</button>
            </div>
          </form>
        </section>
        <section class="panel">
          <h2>Reintentar productos sin imagen</h2>
          {% if retry_runs %}
            <p class="hint">Selecciona una o varias bajadas anteriores. Solo se revisaran las carpetas que todavia contienen <code>SIN_IMAGEN.txt</code>.</p>
            <form method="post" action="{{ url_for('retry_missing') }}">
              <div class="table-tools">
                <button type="button" class="secondary" id="toggle-retry-runs">Marcar todas</button>
              </div>
              <table class="brand-table">
                <thead>
                  <tr>
                    <th>Reintentar</th>
                    <th>Carpeta de descarga</th>
                    <th>Pendientes</th>
                  </tr>
                </thead>
                <tbody>
                  {% for run in retry_runs %}
                    <tr>
                      <td><input class="retry-run-check" type="checkbox" name="retry_run" value="{{ run.name }}" {% if loop.first %}checked{% endif %}></td>
                      <td>{{ run.name }}</td>
                      <td>{{ run.missing_count }}</td>
                    </tr>
                  {% endfor %}
                </tbody>
              </table>
              <div class="grid" style="margin-top: 18px;">
                <div class="span-4">
                  <label for="retry_candidates">Candidatos por producto</label>
                  <input id="retry_candidates" name="retry_candidates" type="number" min="1" max="20" value="8">
                </div>
                <div class="span-4">
                  <label for="retry_serper_api_key">Clave de Serper</label>
                  <input id="retry_serper_api_key" name="serper_api_key" type="password" autocomplete="off" placeholder="{% if serper_available %}Configurada en Windows{% else %}Pegar clave de Serper{% endif %}">
                </div>
                <div class="span-4">
                  <label for="retry_gemini_api_key">Clave de Gemini</label>
                  <input id="retry_gemini_api_key" name="gemini_api_key" type="password" autocomplete="off" placeholder="{% if gemini_available %}Configurada en Windows{% else %}Pegar clave de Gemini{% endif %}">
                </div>
              </div>
              <div class="actions">
                <button type="submit">Reintentar seleccionadas con IA</button>
              </div>
            </form>
            <script>
              (function () {
                const button = document.getElementById('toggle-retry-runs');
                const checks = Array.from(document.querySelectorAll('.retry-run-check'));
                if (!button) return;
                function updateButton() {
                  button.textContent = checks.length && checks.every(function (check) { return check.checked; })
                    ? 'Desmarcar todas' : 'Marcar todas';
                }
                button.addEventListener('click', function () {
                  const shouldCheck = checks.some(function (check) { return !check.checked; });
                  checks.forEach(function (check) { check.checked = shouldCheck; });
                  updateButton();
                });
                checks.forEach(function (check) { check.addEventListener('change', updateButton); });
                updateButton();
              })();
            </script>
          {% else %}
            <p class="hint">No hay carpetas pendientes para reintentar.</p>
          {% endif %}
        </section>
        <section class="panel">
          <h2>Crear imagenes comerciales</h2>
          <p class="hint">Selecciona una carpeta local de origen y otra de destino. Se crea una subcarpeta nueva con los PNG transparentes, sin modificar los archivos originales.</p>
          <form method="post" action="{{ url_for('create_commercial_images') }}">
            <div class="grid">
              <div class="span-12">
                <label for="commercial_source">Carpeta de origen</label>
                <div class="folder-picker">
                  <input id="commercial_source" name="commercial_source" type="text" placeholder="Seleccionar carpeta local" required>
                  <button type="button" class="secondary" data-folder-kind="source">Explorar</button>
                </div>
              </div>
              <div class="span-12">
                <label for="commercial_destination">Carpeta de destino</label>
                <div class="folder-picker">
                  <input id="commercial_destination" name="commercial_destination" type="text" placeholder="Seleccionar carpeta local" required>
                  <button type="button" class="secondary" data-folder-kind="destination">Explorar</button>
                </div>
              </div>
            </div>
            <div class="actions">
              <button type="submit">Crear imagenes comerciales</button>
            </div>
          </form>
          <script>
            (function () {
              const inputs = {
                source: document.getElementById('commercial_source'),
                destination: document.getElementById('commercial_destination')
              };
              document.querySelectorAll('[data-folder-kind]').forEach(function (button) {
                button.addEventListener('click', async function () {
                  const kind = button.dataset.folderKind;
                  button.disabled = true;
                  try {
                    const response = await fetch('/choose-folder/' + kind, {
                      method: 'POST',
                      headers: { 'Content-Type': 'application/json' },
                      body: JSON.stringify({ initial_path: inputs[kind].value })
                    });
                    const result = await response.json();
                    if (result.path) inputs[kind].value = result.path;
                    if (result.error) window.alert(result.error);
                  } catch (error) {
                    window.alert('No se pudo abrir el explorador de carpetas.');
                  } finally {
                    button.disabled = false;
                  }
                });
              });
            })();
          </script>
        </section>
        <section class="panel">
          <h2>Buscar y preparar desde listado</h2>
          <p class="hint">Importa el CSV o Excel de productos sin imagen, elige una carpeta de salida y se generan carpetas con <code>producto_comercial.png</code> listas para cargar.</p>
          <form method="post" action="{{ url_for('prepare_listing_images') }}" enctype="multipart/form-data">
            <div class="grid">
              <div class="span-12">
                <label for="listing_file">Archivo de origen</label>
                <input id="listing_file" name="listing_file" type="file" accept=".csv,.xlsx" required>
              </div>
              <div class="span-12">
                <label for="listing_destination">Carpeta de descarga</label>
                <div class="folder-picker">
                  <input id="listing_destination" name="listing_destination" type="text" placeholder="Seleccionar carpeta local" required>
                  <button type="button" class="secondary" id="browse-listing-destination">Explorar</button>
                </div>
              </div>
              <div class="span-4">
                <label for="listing_candidates">Candidatos por producto</label>
                <input id="listing_candidates" name="listing_candidates" type="number" min="1" max="12" value="6">
              </div>
            </div>
            <div class="actions"><button type="submit">Buscar, descargar y preparar</button></div>
          </form>
          <script>
            (function () {
              const input = document.getElementById('listing_destination');
              const button = document.getElementById('browse-listing-destination');
              button.addEventListener('click', async function () {
                button.disabled = true;
                try {
                  const response = await fetch('/choose-folder/destination', {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify({ initial_path: input.value })
                  });
                  const result = await response.json();
                  if (result.path) input.value = result.path;
                  if (result.error) window.alert(result.error);
                } catch (error) {
                  window.alert('No se pudo abrir el explorador de carpetas.');
                } finally {
                  button.disabled = false;
                }
              });
            })();
          </script>
        </section>
        <section class="panel">
          <h2>Listar productos duplicados</h2>
          <p class="hint">Genera un CSV con los productos repetidos. Compara marca, nombre y presentacion, ignorando guiones y el texto descriptivo posterior.</p>
          <form method="post" action="{{ url_for('create_duplicate_report') }}" enctype="multipart/form-data">
            <div class="grid">
              <div class="span-12">
                <label for="duplicate_file">Archivo de origen</label>
                <input id="duplicate_file" name="duplicate_file" type="file" accept=".csv,.xlsx" required>
              </div>
            </div>
            <div class="actions"><button type="submit">Generar listado de duplicados</button></div>
          </form>
        </section>
        <section class="panel">
          <h2>Cargar imagenes en Pinto Fruta</h2>
          <p class="hint">La carga usa la sesion que abras en Edge. Busca cada producto por su carpeta y carga la imagen comercial: si ya tiene una, la reemplaza.</p>
          <div class="actions">
            <button type="button" class="secondary" id="open-pintofruta">Abrir Pinto Fruta</button>
            <form method="post" action="{{ url_for('create_pintofruta_missing_report') }}" style="display:inline;">
              <button type="submit" class="secondary">Listar productos sin imagen</button>
            </form>
            <form method="post" action="{{ url_for('create_pintofruta_duplicate_report') }}" style="display:inline;">
              <button type="submit" class="secondary">Listar productos duplicados</button>
            </form>
            <form method="post" action="{{ url_for('cleanup_pintofruta_duplicates') }}" style="display:inline;">
              <button type="submit">Eliminar duplicados</button>
            </form>
          </div>
          <form method="post" action="{{ url_for('upload_pintofruta_images') }}">
            <div class="grid">
              <div class="span-12">
                <label for="pintofruta_source">Carpeta de imagenes comerciales</label>
                <div class="folder-picker">
                  <input id="pintofruta_source" name="pintofruta_source" type="text" placeholder="Seleccionar la carpeta que contiene producto_comercial.png" required>
                  <button type="button" class="secondary" id="browse-pintofruta-source">Explorar</button>
                </div>
              </div>
              <div class="span-12">
                <label><input name="pintofruta_ai_matching" type="checkbox" value="1" checked {% if not gemini_available %}disabled{% endif %}> Usar IA para comparar el envase y resolver coincidencias ambiguas</label>
                {% if not gemini_available %}<p class="hint">Configura <code>GEMINI_API_KEY</code> para habilitar este desempate.</p>{% endif %}
              </div>
            </div>
            <div class="actions">
              <button type="submit">Cargar y reemplazar imagenes</button>
            </div>
          </form>
          <script>
            (function () {
              const sourceInput = document.getElementById('pintofruta_source');
              const browseButton = document.getElementById('browse-pintofruta-source');
              const openButton = document.getElementById('open-pintofruta');
              browseButton.addEventListener('click', async function () {
                browseButton.disabled = true;
                try {
                  const response = await fetch('/choose-folder/source', {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify({ initial_path: sourceInput.value })
                  });
                  const result = await response.json();
                  if (result.path) sourceInput.value = result.path;
                  if (result.error) window.alert(result.error);
                } catch (error) {
                  window.alert('No se pudo abrir el explorador de carpetas.');
                } finally {
                  browseButton.disabled = false;
                }
              });
              openButton.addEventListener('click', async function () {
                openButton.disabled = true;
                try {
                  const response = await fetch('/pintofruta/open-browser', { method: 'POST' });
                  const result = await response.json();
                  window.alert(result.message || result.error || 'No se pudo abrir Pinto Fruta.');
                } catch (error) {
                  window.alert('No se pudo abrir Pinto Fruta.');
                } finally {
                  openButton.disabled = false;
                }
              });
            })();
          </script>
        </section>
        <section class="panel">
          <h2>Instagram y disponibilidad</h2>
          <p class="hint">Audita las publicaciones de <code>@distribuidora_pintofruta</code> con Gemini y compara los productos detectados contra el catalogo. El CSV resultante se puede usar para inactivar productos que ya no deben mostrarse, sin borrarlos.</p>
          <div class="actions">
            <form method="post" action="{{ url_for('create_pintofruta_instagram_audit') }}" style="display:inline;">
              <button type="submit" class="secondary">Auditar catalogo contra Instagram</button>
            </form>
          </div>
          <form method="post" action="{{ url_for('deactivate_pintofruta_report') }}" enctype="multipart/form-data" style="margin-top: 18px;">
            <div class="grid">
              <div class="span-12">
                <label for="instagram_deactivation_report">CSV de productos no detectados en Instagram</label>
                <input id="instagram_deactivation_report" name="instagram_deactivation_report" type="file" accept=".csv" required>
              </div>
            </div>
            <div class="actions"><button type="submit">Inactivar productos del CSV</button></div>
          </form>
        </section>
        <section class="panel">
          <h2>Marcas sin productos activos</h2>
          <p class="hint">Revisa directamente la base de Pinto Fruta. La vista previa solo genera un CSV; al inactivar se oculta la marca y se quita de destacados. No borra datos ni requiere que Edge este abierto.</p>
          <form method="post" action="{{ url_for('manage_pintofruta_empty_brands') }}">
            <div class="grid">
              <div class="span-12">
                <label for="pintofruta_env_file">Archivo .env.local de Pinto Fruta</label>
                <div class="folder-picker">
                  <input id="pintofruta_env_file" name="pintofruta_env_file" type="text" value="{{ pintofruta_env_default }}" placeholder="Seleccionar el archivo .env.local" required>
                  <button type="button" class="secondary" id="browse-pintofruta-env">Explorar</button>
                </div>
              </div>
            </div>
            <div class="actions">
              <button type="submit" class="secondary" name="empty_brand_action" value="preview">Vista previa y CSV</button>
              <button type="submit" name="empty_brand_action" value="apply">Inactivar marcas vacias</button>
            </div>
          </form>
          <script>
            (function () {
              const input = document.getElementById('pintofruta_env_file');
              const button = document.getElementById('browse-pintofruta-env');
              button.addEventListener('click', async function () {
                button.disabled = true;
                try {
                  const response = await fetch('/choose-env-file', {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify({ initial_path: input.value })
                  });
                  const result = await response.json();
                  if (result.path) input.value = result.path;
                  if (result.error) window.alert(result.error);
                } catch (error) {
                  window.alert('No se pudo abrir el explorador de archivos.');
                } finally {
                  button.disabled = false;
                }
              });
            })();
          </script>
        </section>
        <section class="panel">
          <h2>Reportes guardados</h2>
          <p class="hint">Los CSV de Pinto Fruta quedan en <code>jobs</code> y siguen disponibles aunque cierres o reinicies la aplicacion.</p>
          {% if pintofruta_reports %}
            <table class="brand-table">
              <thead><tr><th>Reporte</th><th>Generado</th><th>Archivo</th></tr></thead>
              <tbody>
                {% for report in pintofruta_reports %}
                  <tr>
                    <td>{{ report.name }}</td>
                    <td>{{ report.modified }}</td>
                    <td><a class="button secondary" href="{{ url_for('download_saved_pintofruta_report', report_name=report.name) }}">Descargar CSV</a></td>
                  </tr>
                {% endfor %}
              </tbody>
            </table>
          {% else %}
            <p class="hint">Aun no hay reportes de Pinto Fruta guardados.</p>
          {% endif %}
        </section>
        """
        ,
        retry_runs=available_download_runs(),
        serper_available=bool(get_config_value("SERPER_API_KEY")),
        gemini_available=bool(get_config_value("GEMINI_API_KEY")),
        pintofruta_env_default=str(PINFOFRUTA_ENV_DEFAULT) if PINFOFRUTA_ENV_DEFAULT.is_file() else "",
        pintofruta_reports=saved_pintofruta_reports(),
    )


@app.post("/retry-missing")
def retry_missing():
    available_names = {str(run["name"]) for run in available_download_runs()}
    run_names = [name for name in request.form.getlist("retry_run") if name in available_names]
    if not run_names:
        return redirect(url_for("index"))

    serper_api_key = request.form.get("serper_api_key", "").strip() or get_config_value("SERPER_API_KEY")
    gemini_api_key = request.form.get("gemini_api_key", "").strip() or get_config_value("GEMINI_API_KEY")
    if not serper_api_key:
        return page(
            """
            <section class="panel">
              <h2>Falta la clave de Google Images</h2>
              <p>Para reintentar las carpetas pendientes, pega una clave de Serper o guardala en Windows.</p>
              <div class="actions"><a class="button" href="{{ url_for('index') }}">Volver</a></div>
            </section>
            """
        ), 400

    try:
        max_candidates = max(1, min(20, int(request.form.get("retry_candidates", "8"))))
    except ValueError:
        max_candidates = 8

    job_id = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[job_id] = JobStatus(id=job_id)
    thread = threading.Thread(
        target=run_retry_missing_images_job,
        kwargs={
            "job_id": job_id,
            "run_names": run_names,
            "serper_api_key": serper_api_key,
            "gemini_api_key": gemini_api_key,
            "max_candidates": max_candidates,
        },
        daemon=True,
    )
    thread.start()
    return redirect(url_for("job", job_id=job_id))


@app.post("/pintofruta/open-browser")
def open_pintofruta_browser_route():
    try:
        return jsonify(message=open_pintofruta_browser())
    except RuntimeError as exc:
        return jsonify(error=str(exc)), 500


@app.post("/pintofruta-empty-brands")
def manage_pintofruta_empty_brands():
    env_file = request.form.get("pintofruta_env_file", "").strip()
    action = request.form.get("empty_brand_action", "preview")
    if not env_file:
        return redirect(url_for("index"))
    if action not in {"preview", "apply"}:
        return redirect(url_for("index"))

    job_id = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[job_id] = JobStatus(id=job_id)
    thread = threading.Thread(
        target=run_pintofruta_empty_brands_job,
        kwargs={
            "job_id": job_id,
            "env_file": env_file,
            "apply_changes": action == "apply",
        },
        daemon=True,
    )
    thread.start()
    return redirect(url_for("job", job_id=job_id))


@app.post("/pintofruta-instagram-audit")
def create_pintofruta_instagram_audit():
    if not pintofruta_browser_connected():
        return page(
            """
            <section class="panel">
              <h2>Primero abre Pinto Fruta</h2>
              <p>Abre la sesion de Pinto Fruta antes de auditar Instagram.</p>
              <div class="actions"><a class="button" href="{{ url_for('index') }}">Volver</a></div>
            </section>
            """
        ), 400
    api_key = get_config_value("GEMINI_API_KEY")
    if not api_key:
        return page(
            """
            <section class="panel">
              <h2>Falta la clave de Gemini</h2>
              <p>Configura <code>GEMINI_API_KEY</code> en Windows para analizar las publicaciones de Instagram.</p>
              <div class="actions"><a class="button" href="{{ url_for('index') }}">Volver</a></div>
            </section>
            """
        ), 400
    job_id = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[job_id] = JobStatus(id=job_id)
    thread = threading.Thread(target=run_pintofruta_instagram_audit_job, args=(job_id, api_key), daemon=True)
    thread.start()
    return redirect(url_for("job", job_id=job_id))


@app.post("/pintofruta-deactivate-report")
def deactivate_pintofruta_report():
    source = request.files.get("instagram_deactivation_report")
    if not source or Path(source.filename or "").suffix.lower() != ".csv":
        return page(
            """
            <section class="panel">
              <h2>CSV requerido</h2>
              <p>Selecciona el CSV generado por la auditoria de Instagram.</p>
              <div class="actions"><a class="button" href="{{ url_for('index') }}">Volver</a></div>
            </section>
            """
        ), 400
    if not pintofruta_browser_connected():
        return page(
            """
            <section class="panel">
              <h2>Primero abre Pinto Fruta</h2>
              <p>Abre la sesion de Pinto Fruta antes de inactivar productos.</p>
              <div class="actions"><a class="button" href="{{ url_for('index') }}">Volver</a></div>
            </section>
            """
        ), 400
    upload_path = UPLOAD_DIR / f"instagram-audit-{uuid.uuid4().hex}.csv"
    source.save(upload_path)
    job_id = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[job_id] = JobStatus(id=job_id)
    thread = threading.Thread(target=run_pintofruta_deactivate_report_job, args=(job_id, str(upload_path)), daemon=True)
    thread.start()
    return redirect(url_for("job", job_id=job_id))


@app.post("/pintofruta-missing-report")
def create_pintofruta_missing_report():
    if not pintofruta_browser_connected():
        return page(
            """
            <section class="panel">
              <h2>Primero abre Pinto Fruta</h2>
              <p>Abre la sesion de Pinto Fruta antes de generar el listado.</p>
              <div class="actions"><a class="button" href="{{ url_for('index') }}">Volver</a></div>
            </section>
            """
        ), 400
    job_id = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[job_id] = JobStatus(id=job_id)
    thread = threading.Thread(target=run_pintofruta_missing_report_job, args=(job_id,), daemon=True)
    thread.start()
    return redirect(url_for("job", job_id=job_id))


@app.post("/pintofruta-duplicate-report")
def create_pintofruta_duplicate_report():
    if not pintofruta_browser_connected():
        return page(
            """
            <section class="panel">
              <h2>Primero abre Pinto Fruta</h2>
              <p>Abre la sesion de Pinto Fruta antes de generar el listado.</p>
              <div class="actions"><a class="button" href="{{ url_for('index') }}">Volver</a></div>
            </section>
            """
        ), 400
    job_id = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[job_id] = JobStatus(id=job_id)
    thread = threading.Thread(target=run_pintofruta_duplicate_report_job, args=(job_id,), daemon=True)
    thread.start()
    return redirect(url_for("job", job_id=job_id))


@app.post("/pintofruta-cleanup-duplicates")
def cleanup_pintofruta_duplicates():
    if not pintofruta_browser_connected():
        return page(
            """
            <section class="panel">
              <h2>Primero abre Pinto Fruta</h2>
              <p>Abre la sesion de Pinto Fruta antes de eliminar duplicados.</p>
              <div class="actions"><a class="button" href="{{ url_for('index') }}">Volver</a></div>
            </section>
            """
        ), 400
    job_id = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[job_id] = JobStatus(id=job_id)
    thread = threading.Thread(target=run_pintofruta_duplicate_cleanup_job, args=(job_id,), daemon=True)
    thread.start()
    return redirect(url_for("job", job_id=job_id))


@app.post("/pintofruta-upload")
def upload_pintofruta_images():
    source_path = request.form.get("pintofruta_source", "").strip()
    ai_matching = request.form.get("pintofruta_ai_matching") == "1"
    if not source_path:
        return redirect(url_for("index"))
    if not pintofruta_browser_connected():
        return page(
            """
            <section class="panel">
              <h2>Primero abre Pinto Fruta</h2>
              <p>Usa el boton <strong>Abrir Pinto Fruta</strong>, inicia sesion en la ventana de Edge y vuelve a iniciar la carga.</p>
              <div class="actions"><a class="button" href="{{ url_for('index') }}">Volver</a></div>
            </section>
            """
        ), 400

    job_id = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[job_id] = JobStatus(id=job_id)
    thread = threading.Thread(
        target=run_pintofruta_upload_job,
        kwargs={
            "job_id": job_id,
            "source_path": source_path,
            "ai_matching": ai_matching,
        },
        daemon=True,
    )
    thread.start()
    return redirect(url_for("job", job_id=job_id))


@app.get("/job/<job_id>/manual-choice/<choice_id>/image")
def manual_choice_image(job_id: str, choice_id: str):
    status = JOBS.get(job_id)
    if status is None:
        return "", 404
    choice = next((item for item in status.manual_choices if item.id == choice_id), None)
    if choice is None or not choice.source_image.is_file():
        return "", 404
    return send_file(choice.source_image)


@app.post("/pintofruta-manual/<job_id>")
def upload_manual_pintofruta_choices(job_id: str):
    previous_job = JOBS.get(job_id)
    if previous_job is None:
        return redirect(url_for("index"))
    if not pintofruta_browser_connected():
        return page(
            """
            <section class="panel">
              <h2>Primero abre Pinto Fruta</h2>
              <p>Abre la sesion de Pinto Fruta antes de cargar las elecciones manuales.</p>
              <div class="actions"><a class="button" href="{{ url_for('index') }}">Volver</a></div>
            </section>
            """
        ), 400

    selections: list[tuple[ManualChoice, ManualCandidate]] = []
    for choice in previous_job.manual_choices:
        selected = request.form.get(f"choice__{choice.id}")
        try:
            candidate = choice.candidates[int(selected)]
        except (TypeError, ValueError, IndexError):
            continue
        selections.append((choice, candidate))
    if not selections:
        return redirect(url_for("job", job_id=job_id))

    new_job_id = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[new_job_id] = JobStatus(id=new_job_id)
    thread = threading.Thread(
        target=run_manual_pintofruta_upload_job,
        kwargs={"job_id": new_job_id, "selections": selections},
        daemon=True,
    )
    thread.start()
    return redirect(url_for("job", job_id=new_job_id))


@app.post("/choose-folder/<kind>")
def choose_folder(kind: str):
    titles = {
        "source": "Elegir carpeta de origen",
        "destination": "Elegir carpeta de destino",
    }
    if kind not in titles:
        return jsonify(error="Selector de carpeta no valido"), 400

    payload = request.get_json(silent=True) or {}
    initial_path = Path(str(payload.get("initial_path", ""))).expanduser()
    if not initial_path.is_dir():
        initial_path = DOWNLOAD_DIR

    try:
        import tkinter as tk
        from tkinter import filedialog

        with FOLDER_DIALOG_LOCK:
            window = tk.Tk()
            window.withdraw()
            window.attributes("-topmost", True)
            try:
                selected_path = filedialog.askdirectory(
                    parent=window,
                    title=titles[kind],
                    initialdir=str(initial_path),
                    mustexist=True,
                )
            finally:
                window.destroy()
        return jsonify(path=selected_path)
    except Exception:
        return jsonify(error="No se pudo abrir el explorador de carpetas"), 500


@app.post("/choose-env-file")
def choose_env_file():
    payload = request.get_json(silent=True) or {}
    initial_path = Path(str(payload.get("initial_path", ""))).expanduser()
    initial_dir = initial_path.parent if initial_path.is_file() else initial_path
    if not initial_dir.is_dir():
        initial_dir = PINFOFRUTA_ENV_DEFAULT.parent if PINFOFRUTA_ENV_DEFAULT.parent.is_dir() else BASE_DIR

    try:
        import tkinter as tk
        from tkinter import filedialog

        with FOLDER_DIALOG_LOCK:
            window = tk.Tk()
            window.withdraw()
            window.attributes("-topmost", True)
            try:
                selected_path = filedialog.askopenfilename(
                    parent=window,
                    title="Elegir .env.local de Pinto Fruta",
                    initialdir=str(initial_dir),
                    initialfile=initial_path.name if initial_path.is_file() else ".env.local",
                    filetypes=[("Archivos de entorno", ".env .local"), ("Todos los archivos", "*.*")],
                )
            finally:
                window.destroy()
        return jsonify(path=selected_path)
    except Exception:
        return jsonify(error="No se pudo abrir el explorador de archivos"), 500


@app.post("/create-commercial-images")
def create_commercial_images():
    source_path = request.form.get("commercial_source", "").strip()
    destination_path = request.form.get("commercial_destination", "").strip()
    if not source_path or not destination_path:
        return redirect(url_for("index"))

    job_id = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[job_id] = JobStatus(id=job_id)
    thread = threading.Thread(
        target=run_commercial_images_job,
        kwargs={
            "job_id": job_id,
            "source_path": source_path,
            "destination_path": destination_path,
        },
        daemon=True,
    )
    thread.start()
    return redirect(url_for("job", job_id=job_id))


@app.post("/prepare-listing-images")
def prepare_listing_images():
    listing_file = request.files.get("listing_file")
    destination_path = request.form.get("listing_destination", "").strip()
    if not listing_file or not destination_path:
        return redirect(url_for("index"))
    suffix = Path(listing_file.filename or "").suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        return page(
            """
            <section class="panel">
              <h2>Formato no compatible</h2>
              <p>Selecciona un archivo CSV o XLSX.</p>
              <div class="actions"><a class="button" href="{{ url_for('index') }}">Volver</a></div>
            </section>
            """
        ), 400
    try:
        max_candidates = max(1, min(12, int(request.form.get("listing_candidates", "6"))))
    except ValueError:
        max_candidates = 6
    upload_path = UPLOAD_DIR / f"listado-{uuid.uuid4().hex}{suffix}"
    listing_file.save(upload_path)
    job_id = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[job_id] = JobStatus(id=job_id)
    thread = threading.Thread(
        target=run_listing_image_preparation_job,
        kwargs={
            "job_id": job_id,
            "listing_path": str(upload_path),
            "destination_path": destination_path,
            "max_candidates": max_candidates,
        },
        daemon=True,
    )
    thread.start()
    return redirect(url_for("job", job_id=job_id))


@app.post("/duplicate-report")
def create_duplicate_report():
    duplicate_file = request.files.get("duplicate_file")
    if not duplicate_file:
        return redirect(url_for("index"))
    suffix = Path(duplicate_file.filename or "").suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        return page(
            """
            <section class="panel">
              <h2>Formato no compatible</h2>
              <p>Selecciona un archivo CSV o XLSX.</p>
              <div class="actions"><a class="button" href="{{ url_for('index') }}">Volver</a></div>
            </section>
            """
        ), 400

    upload_path = UPLOAD_DIR / f"duplicados-{uuid.uuid4().hex}{suffix}"
    duplicate_file.save(upload_path)
    job_id = uuid.uuid4().hex
    with JOBS_LOCK:
        JOBS[job_id] = JobStatus(id=job_id)
    thread = threading.Thread(
        target=run_duplicate_listing_job,
        kwargs={"job_id": job_id, "listing_path": str(upload_path)},
        daemon=True,
    )
    thread.start()
    return redirect(url_for("job", job_id=job_id))


@app.post("/upload")
def upload():
    file = request.files.get("excel")
    if not file or not file.filename.lower().endswith(".xlsx"):
        return redirect(url_for("index"))

    upload_id = uuid.uuid4().hex
    excel_path = UPLOAD_DIR / f"{upload_id}.xlsx"
    file.save(excel_path)
    layout = read_excel_layout(excel_path)
    return page(
        """
        <section class="panel">
          <h2>2. Elegir columnas</h2>
          <p class="hint">Encabezados detectados en la fila {{ layout.header_row }}.</p>
          <form method="post" action="{{ url_for('configure') }}">
            <input type="hidden" name="upload_id" value="{{ upload_id }}">
            <input type="hidden" name="header_row" value="{{ layout.header_row }}">
            <div class="grid">
              <div class="span-3">
                <label for="brand_column">Columna de marca</label>
                <select id="brand_column" name="brand_column" required>
                  {% for header in layout.headers %}
                    <option value="{{ header }}" {% if 'marca' in header.lower() or 'brand' in header.lower() %}selected{% endif %}>{{ header }}</option>
                  {% endfor %}
                </select>
              </div>
              <div class="span-3">
                <label for="product_column">Columna de producto</label>
                <select id="product_column" name="product_column" required>
                  {% for header in layout.headers %}
                    <option value="{{ header }}" {% if 'producto' in header.lower() or 'descripcion' in header.lower() or 'nombre' in header.lower() or 'detalle' in header.lower() or 'variedad' in header.lower() %}selected{% endif %}>{{ header }}</option>
                  {% endfor %}
                </select>
              </div>
              <div class="span-3">
                <label for="extra_column">Columna extra opcional</label>
                <select id="extra_column" name="extra_column">
                  <option value="">Sin extra</option>
                  {% for header in layout.headers %}
                    <option value="{{ header }}" {% if 'presentacion' in header.lower() or 'medida' in header.lower() or 'tamano' in header.lower() %}selected{% endif %}>{{ header }}</option>
                  {% endfor %}
                </select>
              </div>
              <div class="span-3">
                <label for="code_column">Codigo/SKU opcional</label>
                <select id="code_column" name="code_column">
                  <option value="">Sin codigo</option>
                  {% for header in layout.headers %}
                    <option value="{{ header }}" {% if 'sku' in header.lower() or 'codigo' in header.lower() or 'cod' == header.lower() %}selected{% endif %}>{{ header }}</option>
                  {% endfor %}
                </select>
              </div>
            </div>
            <div class="actions">
              <button type="submit">Detectar marcas</button>
            </div>
          </form>
        </section>
        """,
        upload_id=upload_id,
        layout=layout,
    )


@app.post("/configure")
def configure():
    upload_id = request.form["upload_id"]
    excel_path = UPLOAD_DIR / f"{upload_id}.xlsx"
    header_row = int(request.form.get("header_row", "1"))
    brand_column = request.form["brand_column"]
    product_column = request.form["product_column"]
    code_column = request.form.get("code_column", "")
    extra_column = request.form.get("extra_column", "")
    products = parse_excel(excel_path, header_row, brand_column, product_column, code_column, extra_column)
    groups = grouped_brands(products)

    return page(
        """
        <section class="panel">
          <h2>3. Web por marca</h2>
          <form method="post" action="{{ url_for('start') }}">
            <input type="hidden" name="upload_id" value="{{ upload_id }}">
            <input type="hidden" name="header_row" value="{{ header_row }}">
            <input type="hidden" name="brand_column" value="{{ brand_column }}">
          <input type="hidden" name="product_column" value="{{ product_column }}">
          <input type="hidden" name="code_column" value="{{ code_column }}">
          <input type="hidden" name="extra_column" value="{{ extra_column }}">
            <div class="table-tools">
              <button type="button" class="secondary" id="toggle-brands">Desmarcar todos</button>
            </div>
            <table class="brand-table">
              <thead>
                <tr>
                  <th>Descargar</th>
                  <th>Marca</th>
                  <th>Productos</th>
                  <th>Web del fabricante</th>
                </tr>
              </thead>
              <tbody>
                {% for brand, items in groups.items() %}
                  <tr>
                    <td>
                      <input class="brand-check" type="checkbox" name="brand_enabled__{{ loop.index0 }}" checked>
                    </td>
                    <td>{{ brand }}</td>
                    <td>{{ items|length }}</td>
                    <td>
                      <input name="brand_url__{{ loop.index0 }}" placeholder="https://fabricante.com/catalogo">
                      <input type="hidden" name="brand_name__{{ loop.index0 }}" value="{{ brand }}">
                    </td>
                  </tr>
                {% endfor %}
              </tbody>
            </table>
            <input type="hidden" name="brand_count" value="{{ groups|length }}">
            <div class="grid" style="margin-top: 18px;">
              <div class="span-4">
                <label for="max_pages">Paginas max. por marca</label>
                <input id="max_pages" name="max_pages" type="number" min="1" max="100" value="25">
              </div>
              <div class="span-4">
                <label for="max_images">Imagenes por producto</label>
                <input id="max_images" name="max_images" type="number" min="1" max="10" value="3">
              </div>
              <div class="span-4">
                <label for="min_score">Puntaje minimo</label>
                <input id="min_score" name="min_score" type="number" min="1" max="50" value="8">
              </div>
            </div>
            <p class="hint">Mas paginas aumenta posibilidades y tambien demora. Si encuentra pocas imagenes, bajar el puntaje minimo a 4 o 5.</p>
            <div class="grid" style="margin-top: 18px;">
              <div class="span-6">
                <label for="serper_api_key">Clave de Serper para Google Images</label>
                <input id="serper_api_key" name="serper_api_key" type="password" autocomplete="off" placeholder="{% if serper_available %}Configurada en Windows{% else %}Pegar clave de Serper{% endif %}">
                <p class="hint">Necesaria para los botones de Internet e IA. <a href="https://serper.dev/" target="_blank" rel="noopener noreferrer">Obtener clave gratuita</a>.</p>
              </div>
            </div>
            <div class="actions">
              <button type="submit" name="internet_search" value="0">Iniciar descarga</button>
              <button type="submit" class="secondary" name="internet_search" value="1">Buscar en Google Images</button>
              <button type="submit" name="ai_validation" value="1">Buscar y validar con IA</button>
            </div>
            <p class="hint">La busqueda usa Google Images. La validacion con IA tambien comprueba el envase con Gemini.</p>
          </form>
        </section>
        <script>
          (function () {
            const button = document.getElementById('toggle-brands');
            const checks = Array.from(document.querySelectorAll('.brand-check'));
            if (!button) return;

            function updateButton() {
              const allChecked = checks.length > 0 && checks.every(function (check) {
                return check.checked;
              });
              button.textContent = allChecked ? 'Desmarcar todos' : 'Marcar todos';
            }

            button.addEventListener('click', function () {
              const shouldCheck = checks.some(function (check) {
                return !check.checked;
              });
              checks.forEach(function (check) {
                check.checked = shouldCheck;
              });
              updateButton();
            });

            checks.forEach(function (check) {
              check.addEventListener('change', updateButton);
            });
            updateButton();
          })();
        </script>
        """,
        upload_id=upload_id,
        header_row=header_row,
        brand_column=brand_column,
        product_column=product_column,
        code_column=code_column,
        extra_column=extra_column,
        groups=groups,
        serper_available=bool(get_config_value("SERPER_API_KEY")),
    )


@app.post("/start")
def start():
    upload_id = request.form["upload_id"]
    excel_path = UPLOAD_DIR / f"{upload_id}.xlsx"
    brand_count = int(request.form.get("brand_count", "0"))
    brand_urls = {}
    brand_enabled = {}
    for index in range(brand_count):
        brand = request.form.get(f"brand_name__{index}", "")
        url = request.form.get(f"brand_url__{index}", "")
        if brand:
            brand_urls[brand] = safe_url(url) if url else ""
            brand_enabled[brand] = request.form.get(f"brand_enabled__{index}") == "on"
    ai_validation = request.form.get("ai_validation") == "1"
    internet_search = request.form.get("internet_search") == "1" or ai_validation
    serper_api_key = request.form.get("serper_api_key", "").strip() or get_config_value("SERPER_API_KEY")
    if internet_search and not serper_api_key:
        return page(
            """
            <section class="panel">
              <h2>Falta la clave de Google Images</h2>
              <p>Para buscar imagenes automaticamente, pega una clave de Serper en el campo de la pantalla anterior.</p>
              <p><a href="https://serper.dev/" target="_blank" rel="noopener noreferrer">Crear clave gratuita en Serper</a></p>
              <div class="actions"><button type="button" onclick="history.back()">Volver</button></div>
            </section>
            """
        ), 400

    job_id = uuid.uuid4().hex
    status = JobStatus(id=job_id)
    with JOBS_LOCK:
        JOBS[job_id] = status

    thread = threading.Thread(
        target=run_download_job,
        kwargs={
            "job_id": job_id,
            "excel_path": excel_path,
            "header_row": int(request.form.get("header_row", "1")),
            "brand_column": request.form["brand_column"],
            "product_column": request.form["product_column"],
            "code_column": request.form.get("code_column", ""),
            "extra_column": request.form.get("extra_column", ""),
            "brand_urls": brand_urls,
            "brand_enabled": brand_enabled,
            "internet_search": internet_search,
            "ai_validation": ai_validation,
            "serper_api_key": serper_api_key,
            "max_pages": int(request.form.get("max_pages", "25")),
            "max_images": int(request.form.get("max_images", "3")),
            "min_score": int(request.form.get("min_score", "8")),
        },
        daemon=True,
    )
    thread.start()
    return redirect(url_for("job", job_id=job_id))


@app.get("/job/<job_id>")
def job(job_id: str):
    status = JOBS.get(job_id)
    if status is None:
        return redirect(url_for("index"))

    return page(
        """
        <section class="panel">
          <h2>4. Progreso</h2>
          <p><span class="status {{ status.state }}">{{ status.state }}</span> {{ status.message }}</p>
          <div class="progress" aria-label="Progreso">
            <div class="bar" style="--value: {{ status.progress }}%;"></div>
          </div>
          <div class="actions">
            {% if status.state in ['running', 'queued'] %}
              <a class="button" href="{{ url_for('job', job_id=status.id) }}">Actualizar</a>
            {% endif %}
            {% if status.zip_path %}
              <a class="button" href="{{ url_for('download_zip', job_id=status.id) }}">{% if status.zip_path.suffix.lower() == '.csv' %}Descargar CSV{% else %}Descargar ZIP{% endif %}</a>
            {% endif %}
          </div>
          {% if status.missing_searches %}
            <h3>Revisar en Google Images</h3>
            <p class="hint">Estos productos no tuvieron una descarga valida. Cada boton abre la busqueda exacta en una pestaña nueva.</p>
            <div class="manual-searches">
              {% for label, search_url in status.missing_searches %}
                <a class="button secondary manual-search" href="{{ search_url }}" target="_blank" rel="noopener noreferrer">Buscar {{ label }}</a>
              {% endfor %}
            </div>
          {% endif %}
          {% if status.manual_choices and status.state not in ['running', 'queued'] %}
            <h3>Elegir coincidencias pendientes</h3>
            <p class="hint">La carga se realizara solo para las fichas que elijas aqui.</p>
            <form method="post" action="{{ url_for('upload_manual_pintofruta_choices', job_id=status.id) }}">
              {% for choice in status.manual_choices %}
                <div style="margin: 16px 0; padding: 16px; border-top: 1px solid var(--line);">
                  <div class="grid">
                    <div class="span-3">
                      <img src="{{ url_for('manual_choice_image', job_id=status.id, choice_id=choice.id) }}" alt="{{ choice.product_name }}" style="display:block; width:100%; max-height:180px; object-fit:contain;">
                    </div>
                    <div class="span-9">
                      <strong>{{ choice.product_name }}</strong>
                      <p class="hint">Marca: {{ choice.brand }}</p>
                      {% for candidate in choice.candidates %}
                        <label style="display:block; margin: 10px 0;">
                          <input type="radio" name="choice__{{ choice.id }}" value="{{ loop.index0 }}" required>
                          {{ candidate.title }}
                        </label>
                      {% endfor %}
                    </div>
                  </div>
                </div>
              {% endfor %}
              <div class="actions"><button type="submit">Cargar las elecciones</button></div>
            </form>
          {% endif %}
          {% if status.output_dir %}
            <p class="hint">Carpeta: {{ status.output_dir }}</p>
          {% endif %}
          <pre>{{ "\\n".join(status.log) }}</pre>
        </section>
        {% if status.state in ['running', 'queued'] %}
          <script>setTimeout(function(){ location.reload(); }, 2500);</script>
        {% endif %}
        """,
        status=status,
    )


@app.get("/download/<job_id>")
def download_zip(job_id: str):
    status = JOBS.get(job_id)
    if status is None or status.zip_path is None:
        return redirect(url_for("index"))
    return send_file(status.zip_path, as_attachment=True, download_name=status.zip_path.name)


@app.get("/pintofruta-report/<report_name>")
def download_saved_pintofruta_report(report_name: str):
    path = (JOB_DIR / report_name).resolve()
    allowed_root = JOB_DIR.resolve()
    prefixes = (
        "pintofruta-",
        "catalogo-no-detectado-instagram-",
        "productos-inactivados-instagram-",
    )
    if (
        path.parent != allowed_root
        or path.suffix.lower() != ".csv"
        or not path.name.startswith(prefixes)
        or not path.is_file()
    ):
        return redirect(url_for("index"))
    return send_file(path, as_attachment=True, download_name=path.name)


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    app.run(host="127.0.0.1", port=port, debug=False)
